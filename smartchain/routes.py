import csv
import json
import math
import uuid
from datetime import datetime
from io import BytesIO, StringIO

import os
from flask import Blueprint, Response, abort, current_app, flash, redirect, render_template, request, send_from_directory, session, url_for
from flask_login import current_user, login_required, login_user, logout_user
from sqlalchemy import or_
from sqlalchemy.exc import IntegrityError
from werkzeug.utils import secure_filename

from .extensions import db
from .utils import safe_int
from .models import (
    AnnualProcurementPlan,
    ApprovalRecord,
    ArchiveRecord,
    Asset,
    AuditLog,
    BackupRecord,
    Beneficiary,
    BidAdjudicationCommittee,
    BidEvaluationCommittee,
    BudgetForecast,
    ComplianceCheck,
    Contract,
    DemandRequest,
    Department,
    Directorate,
    Document,
    Employee,
    EvaluatorScore,
    EvaluatorScoreRevision,
    BidSubmission,
    InternalControl,
    IntegrationConnection,
    IntegrationTransaction,
    Notification,
    PerformanceIndicator,
    PasswordPolicy,
    Policy,
    PortfolioProject,
    ProcurementRequest,
    ProcurementIntakeRequest,
    ProcurementRequestDocument,
    ProcurementRequestHistory,
    ProcurementQuotation,
    ProcurementChecklistItem,
    ProcurementCaseEvent,
    ProcurementApproval,
    ProcurementOrder,
    ProcurementDelivery,
    Requisition,
    RiskRegister,
    Role,
    SecurityEvent,
    ServiceFeedback,
    StrategicPlan,
    Stakeholder,
    Supplier,
    Tender,
    TenderBidder,
    TenderBidderProfile,
    TenderComplianceItem,
    TenderCriterion,
    TenderCriterionConfig,
    TenderProfile,
    TenderSubcriterion,
    TenderEvaluatorAssignment,
    TenderScreeningRequirement,
    TenderScreeningResult,
    EvaluatorSubcriterionScore,
    TenderDocument,
    TenderStageEvent,
    TenderCommitteeMeeting,
    TenderCommitteeMember,
    TenderOutcome,
    TenderReportRecord,
    TrainingRecord,
    TripRequest,
    Vehicle,
    Driver,
    TransportAllocation,
    TransportCaseEvent,
    TransportDocument,
    TransportAssignment,
    TransportMaintenance,
    TransportIncident,
    TransportTripAuthorisation,
    TransportVehicleIssue,
    TransportPortalRequest,
    TravelRequest,
    Unit,
    User,
)


def log_audit(module, action, entity, details, user=None):
    db.session.add(
        AuditLog(
            module=module,
            action=action,
            entity=entity,
            details=details,
            user_id=user.id if user else None,
            created_at=datetime.utcnow(),
        )
    )
    db.session.commit()


SCM_WORKFLOW_ROLES = {
    "System Administrator",
    "SCM Manager",
    "SCM Officer / Practitioner",
    "SCM Intake Officer",
    "SCM Clerk",
    "SCM Supervisor",
}

DIRECTORATE_RQ_ROLES = {"System Administrator", "Directorate DD", "Directorate Director", "Directorate Chief Director"}

PROCUREMENT_TRANSITIONS = {
    "Submitted": {"Received by SCM", "Documents Required", "Rejected"},
    "Received by SCM": {"Preliminary Review", "Documents Required", "Rejected"},
    "Preliminary Review": {"Ready for Allocation", "Documents Required", "Rejected"},
    "Documents Required": {"Preliminary Review", "Rejected"},
    "Ready for Allocation": {"Allocated"},
    "Allocated": {"Sourcing in Progress", "Internal Control Review", "Rejected"},
    "Sourcing in Progress": {"Comparative Analysis", "Documents Required", "Internal Control Review"},
    "Comparative Analysis": {"Internal Control Review", "Corrections Required"},
    "Internal Control Review": {"Corrections Required", "Approval", "Completed"},
    "Corrections Required": {"Preliminary Review", "Sourcing in Progress", "Internal Control Review"},
    "Approval": {"Completed", "Corrections Required", "Rejected"},
    "Completed": {"Closed"},
    "Closed": set(),
    "Rejected": set(),
}


def require_role(roles):
    return current_user.is_authenticated and current_user.role_obj and current_user.role_obj.name in roles


def add_request_history(request_record, action, to_status=None, comment=None, from_status=None):
    request_record.history.append(ProcurementRequestHistory(action=action, from_status=from_status, to_status=to_status, comment=comment, user=current_user))


def record_procurement_event(request_record, action, to_status=None, comment=None, from_status=None):
    db.session.add(ProcurementCaseEvent(
        request=request_record,
        action=action,
        from_status=from_status if from_status is not None else request_record.status,
        to_status=to_status or request_record.status,
        comment=comment,
        user=current_user,
        role_name=current_user.role_obj.name if current_user.role_obj else None,
    ))


def require_directorate_rq_role():
    if not current_user.role_obj or current_user.role_obj.name not in DIRECTORATE_RQ_ROLES:
        abort(403)
    if current_user.role_obj.name != "System Administrator" and not current_user.employee:
        abort(403)


def notify_user(user_id, title, message):
    db.session.add(Notification(user_id=user_id, title=title, message=message))


# Authentication

auth_bp = Blueprint("auth", __name__, url_prefix="/auth")


@auth_bp.route("/login", methods=["GET", "POST"])
def login():
    if current_user.is_authenticated and request.method == "GET":
        return redirect(url_for("dashboard.index"))

    if request.method == "POST":
        if current_user.is_authenticated:
            logout_user()
        username = request.form.get("username", "").strip()
        user = User.query.filter(or_(User.username == username, User.email == username)).first()
        if not user:
            employee = Employee.query.filter_by(employee_number=username).first()
            user = employee.user if employee else None
        if user and user.is_active:
            login_user(user)
            user.last_login_at = datetime.utcnow()
            log_audit("Auth", "login", "user", f"User {user.username} signed in.", user=user)
            return redirect(url_for("dashboard.index"))
        flash("Invalid or inactive account.")

    return render_template("login.html")


@auth_bp.route("/register", methods=["GET", "POST"])
def register():
    if current_user.is_authenticated:
        return redirect(url_for("dashboard.index"))
    departments = Department.query.order_by(Department.name).all()
    directorates = Directorate.query.order_by(Directorate.name).all()
    units = Unit.query.order_by(Unit.name).all()
    if request.method == "POST":
        email = request.form.get("email", "").strip().lower()
        employee_number = request.form.get("employee_number", "").strip()
        if User.query.filter_by(email=email).first():
            flash("An account with this email address already exists.")
        elif Employee.query.filter_by(employee_number=employee_number).first():
            flash("An employee with this staff number already exists.")
        else:
            department_id = safe_int(request.form.get("department_id"))
            directorate_id = safe_int(request.form.get("directorate_id"))
            unit_id = safe_int(request.form.get("unit_id"))
            employee_role = Role.query.filter_by(name="Employee").first()
            if None in (department_id, directorate_id, unit_id) or not employee_role:
                flash("Select a valid department, directorate, and unit.")
            else:
                employee = Employee(employee_number=employee_number, first_name=request.form.get("first_name"), last_name=request.form.get("surname"), title=request.form.get("job_title"), position=request.form.get("job_title"), email=email, phone=request.form.get("contact_number"), department_id=department_id, directorate_id=directorate_id, unit_id=unit_id, role_id=employee_role.id, account_status="active")
                db.session.add(employee)
                db.session.flush()
                user = User(username=employee_number, email=email, full_name=f"{employee.first_name} {employee.last_name}", role_id=employee_role.id, employee_id=employee.id)
                db.session.add(user)
                db.session.commit()
                log_audit("Auth", "register", "user", f"Registered employee account {user.username}.", user=user)
                flash("Registration successful. You can now sign in.")
                return redirect(url_for("auth.login"))
    return render_template("register.html", departments=departments, directorates=directorates, units=units)


@auth_bp.route("/change-password", methods=["GET", "POST"])
@login_required
def change_password():
    if request.method == "POST":
        flash("Password management is disabled in this environment.")
        return redirect(url_for("dashboard.index"))
    return render_template("change_password.html")


@auth_bp.route("/forgot-password", methods=["GET", "POST"])
def forgot_password():
    if request.method == "POST":
        email = request.form.get("email", "").strip().lower()
        user = User.query.filter_by(email=email).first()
        if user:
            log_audit("Auth", "password_reset_requested", "user", f"Password reset requested for {user.username}.", user=user)
        flash("Password recovery is disabled because passwordless sign-in is active.")
        return redirect(url_for("auth.login"))
    return render_template("forgot_password.html")


procurement_portal_bp = Blueprint("procurement_portal", __name__, url_prefix="/procurement-request")


def next_procurement_request_number():
    return f"RQ-{datetime.utcnow().year}-{ProcurementIntakeRequest.query.count() + 1:06d}"


@procurement_portal_bp.route("/", methods=["GET", "POST"])
@login_required
def submit_procurement_request():
    require_directorate_rq_role()
    employee = current_user.employee
    if request.method == "POST":
        intake_request = ProcurementIntakeRequest(
            request_number=next_procurement_request_number(),
            requester=current_user,
            employee_name=request.form.get("employee_name") or current_user.full_name,
            employee_number=request.form.get("employee_number") or (employee.employee_number if employee else current_user.username),
            email=request.form.get("email") or current_user.email,
            contact_number=request.form.get("contact_number") or (employee.phone if employee else None),
            directorate=request.form.get("directorate") or (employee.directorate.name if employee and employee.directorate else None),
            chief_directorate=request.form.get("chief_directorate"),
            unit=request.form.get("unit") or (employee.unit.name if employee and employee.unit else None),
            job_title=request.form.get("job_title") or (employee.position if employee else None),
            cost_centre=request.form.get("cost_centre"),
            required_by_date=request.form.get("required_by_date"),
            procurement_category=request.form.get("procurement_category"),
            procurement_type=request.form.get("procurement_type", "Goods"),
            description=request.form.get("description"),
            specification=request.form.get("specification"),
            quantity=safe_int(request.form.get("quantity"), 1),
            estimated_value=float(request.form.get("estimated_value", 0) or 0),
            business_need=request.form.get("business_need"),
            justification=request.form.get("justification"),
            delivery_location=request.form.get("delivery_location"),
            responsible_official=request.form.get("responsible_official"),
            status="Submitted",
        )
        db.session.add(intake_request)
        db.session.flush()
        add_request_history(intake_request, "Submitted", to_status="Submitted", comment="Employee procurement request submitted.")
        record_procurement_event(intake_request, "Directorate RQ submitted to Central SCM", "Submitted", "Submitted by an authorised directorate official.", "Draft")
        for name, reason in [("RQ details", "Core request fields"), ("Business motivation", "Business need and justification"), ("Specifications", "Goods/services specification"), ("Supporting documents", "Configured submission evidence")]:
            db.session.add(ProcurementChecklistItem(request=intake_request, name=name, required=True, status="PENDING", reason=reason))
        for scm_user in User.query.join(Role).filter(Role.name.in_(SCM_WORKFLOW_ROLES)).all():
            notify_user(scm_user.id, "New procurement intake request", f"{intake_request.request_number} is awaiting SCM intake review.")
        db.session.commit()
        log_audit("Procurement Intake", "submit", "procurement_intake_request", f"Submitted {intake_request.request_number}.", user=current_user)
        flash(f"Procurement request {intake_request.request_number} submitted successfully.")
        return redirect(url_for("procurement_portal.my_requests"))
    return render_template("procurement_request_form.html", employee=employee)


@procurement_portal_bp.route("/mine")
@login_required
def my_requests():
    require_directorate_rq_role()
    if current_user.role_obj.name == "System Administrator":
        items = ProcurementIntakeRequest.query.order_by(ProcurementIntakeRequest.created_at.desc()).all()
    else:
        directorate = current_user.employee.directorate.name if current_user.employee and current_user.employee.directorate else None
        items = ProcurementIntakeRequest.query.filter_by(directorate=directorate).order_by(ProcurementIntakeRequest.created_at.desc()).all()
    return render_template("procurement_request_list.html", items=items)


@procurement_portal_bp.route("/<int:request_id>")
@login_required
def request_detail(request_id):
    require_directorate_rq_role()
    query = ProcurementIntakeRequest.query.filter_by(id=request_id)
    if current_user.role_obj.name != "System Administrator":
        directorate = current_user.employee.directorate.name if current_user.employee and current_user.employee.directorate else None
        query = query.filter_by(directorate=directorate)
    intake_request = query.first_or_404()
    return render_template("procurement_request_detail.html", item=intake_request, employee_view=True)


@procurement_portal_bp.route("/<int:request_id>/documents", methods=["POST"])
@login_required
def upload_request_document(request_id):
    require_directorate_rq_role()
    query = ProcurementIntakeRequest.query.filter_by(id=request_id)
    if current_user.role_obj.name != "System Administrator":
        directorate = current_user.employee.directorate.name if current_user.employee and current_user.employee.directorate else None
        query = query.filter_by(directorate=directorate)
    intake_request = query.first_or_404()
    upload = request.files.get("document")
    if not upload or not upload.filename:
        flash("Select a document to upload.")
        return redirect(url_for("procurement_portal.request_detail", request_id=request_id))
    filename = secure_filename(upload.filename)
    upload_dir = os.path.join(current_app.instance_path, "procurement_uploads", intake_request.request_number)
    os.makedirs(upload_dir, exist_ok=True)
    storage_path = os.path.join(upload_dir, filename)
    upload.save(storage_path)
    intake_request.documents.append(ProcurementRequestDocument(category=request.form.get("category", "Supporting Document"), title=request.form.get("title") or filename, file_name=filename, storage_path=storage_path, uploaded_by=current_user))
    add_request_history(intake_request, "Document uploaded", comment=f"Requester uploaded {filename}.")
    record_procurement_event(intake_request, "Document uploaded", intake_request.status, f"Uploaded {filename}.")
    db.session.commit()
    log_audit("Procurement Intake", "upload", "procurement_request_document", f"Uploaded {filename} to {intake_request.request_number}.", user=current_user)
    flash("Supporting document uploaded successfully.")
    return redirect(url_for("procurement_portal.request_detail", request_id=request_id))


procurement_intake_bp = Blueprint("procurement_intake", __name__, url_prefix="/procurement-intake")


@procurement_intake_bp.before_request
def protect_procurement_intake():
    if not require_role(SCM_WORKFLOW_ROLES):
        abort(403)


@procurement_intake_bp.route("/")
def queue():
    status = request.args.get("status")
    query = ProcurementIntakeRequest.query
    if status:
        query = query.filter_by(status=status)
    items = query.order_by(ProcurementIntakeRequest.created_at.asc()).all()
    return render_template("procurement_command.html", items=items, active_filter=status, now=datetime.utcnow(), stats={
        "open": ProcurementIntakeRequest.query.filter(~ProcurementIntakeRequest.status.in_(["Closed", "Rejected"])).count(),
        "submitted": ProcurementIntakeRequest.query.filter(ProcurementIntakeRequest.status.in_(["Submitted", "Received by SCM"])).count(),
        "sourcing": ProcurementIntakeRequest.query.filter(ProcurementIntakeRequest.status.in_(["Sourcing in Progress", "Comparative Analysis"])).count(),
        "approval": ProcurementIntakeRequest.query.filter(ProcurementIntakeRequest.status.in_(["Approval", "Internal Control Review"])).count(),
        "delivery": ProcurementIntakeRequest.query.filter(ProcurementIntakeRequest.status.in_(["Completed", "Closed"])).count(),
        "overdue": ProcurementIntakeRequest.query.filter(ProcurementIntakeRequest.required_by_date < datetime.utcnow().strftime("%Y-%m-%d"), ~ProcurementIntakeRequest.status.in_(["Closed", "Rejected"])).count(),
    })


@procurement_intake_bp.route("/<int:request_id>", methods=["GET", "POST"])
def intake_detail(request_id):
    intake_request = ProcurementIntakeRequest.query.get_or_404(request_id)
    if request.method == "POST":
        previous_status = intake_request.status
        new_status = request.form.get("status") or previous_status
        if new_status != previous_status and new_status not in PROCUREMENT_TRANSITIONS.get(previous_status, set()):
            flash(f"Invalid procurement transition from {previous_status} to {new_status}.")
            return redirect(url_for("procurement_intake.intake_detail", request_id=request_id)), 409
        intake_request.status = new_status
        intake_request.scm_notes = request.form.get("scm_notes")
        assigned_id = safe_int(request.form.get("assigned_user_id"))
        if assigned_id:
            intake_request.assigned_user_id = assigned_id
        add_request_history(intake_request, "SCM review", from_status=previous_status, to_status=new_status, comment=intake_request.scm_notes)
        record_procurement_event(intake_request, "SCM status update", new_status, intake_request.scm_notes, previous_status)
        if assigned_id:
            record_procurement_event(intake_request, "Case assigned", new_status, f"Assigned to {intake_request.assigned_user.full_name}.")
        notify_user(intake_request.requester_user_id, f"Procurement request {intake_request.request_number} updated", f"Your request status is now {new_status}.")
        db.session.commit()
        log_audit("Procurement Intake", "review", "procurement_intake_request", f"Updated {intake_request.request_number} from {previous_status} to {new_status}.", user=current_user)
        flash("Procurement intake update saved.")
        return redirect(url_for("procurement_intake.intake_detail", request_id=request_id))
    officers = User.query.join(Role).filter(Role.name.in_(SCM_WORKFLOW_ROLES)).order_by(User.full_name).all()
    return render_template("procurement_request_detail.html", item=intake_request, employee_view=False, officers=officers)


@procurement_intake_bp.route("/<int:request_id>/quotations", methods=["POST"])
def add_quotation(request_id):
    intake_request = ProcurementIntakeRequest.query.get_or_404(request_id)
    quotation = ProcurementQuotation(request=intake_request, supplier_name=request.form.get("supplier_name"), quotation_date=request.form.get("quotation_date"), amount=float(request.form.get("amount", 0) or 0), response_status=request.form.get("response_status", "Received"), compliance_notes=request.form.get("compliance_notes"), created_by=current_user)
    intake_request.status = "Comparative Analysis"
    add_request_history(intake_request, "Quotation captured", to_status=intake_request.status, comment=f"Captured quotation from {quotation.supplier_name}.")
    record_procurement_event(intake_request, "Quotation captured", intake_request.status, f"Captured quotation from {quotation.supplier_name}.")
    db.session.add(quotation)
    db.session.commit()
    log_audit("Procurement Intake", "quotation", "procurement_quotation", f"Captured quotation for {intake_request.request_number}.", user=current_user)
    flash("Quotation captured successfully.")
    return redirect(url_for("procurement_intake.intake_detail", request_id=request_id))


@procurement_intake_bp.route("/<int:request_id>/checklist", methods=["POST"])
def review_checklist(request_id):
    intake_request = ProcurementIntakeRequest.query.get_or_404(request_id)
    item = ProcurementChecklistItem.query.filter_by(id=safe_int(request.form.get("item_id")), request_id=request_id).first_or_404()
    result = request.form.get("status", "PENDING")
    if result not in {"COMPLETE", "INCOMPLETE", "REQUIRES CLARIFICATION", "PENDING"}:
        abort(400)
    item.status = result
    item.evidence_reference = request.form.get("evidence_reference")
    item.reason = request.form.get("reason") or item.reason
    item.reviewed_by = current_user
    item.reviewed_at = datetime.utcnow()
    if result in {"INCOMPLETE", "REQUIRES CLARIFICATION"}:
        intake_request.status = "Documents Required"
    record_procurement_event(intake_request, "Completeness review", intake_request.status, f"{item.name}: {result}.")
    db.session.commit()
    flash("Checklist item updated.")
    return redirect(url_for("procurement_intake.intake_detail", request_id=request_id))


@procurement_intake_bp.route("/<int:request_id>/approval", methods=["POST"])
def decide_procurement_approval(request_id):
    intake_request = ProcurementIntakeRequest.query.get_or_404(request_id)
    decision = request.form.get("status", "PENDING")
    if decision not in {"APPROVED", "REJECTED", "RETURNED_FOR_CORRECTION"}:
        abort(400)
    approval = ProcurementApproval(request=intake_request, status=decision, decision=request.form.get("decision"), approver=current_user, decided_at=datetime.utcnow())
    db.session.add(approval)
    target = "Completed" if decision == "APPROVED" else "Corrections Required" if decision == "RETURNED_FOR_CORRECTION" else "Rejected"
    previous = intake_request.status
    intake_request.status = target
    add_request_history(intake_request, "Procurement approval", from_status=previous, to_status=target, comment=approval.decision)
    record_procurement_event(intake_request, "Approval decision", target, f"{decision}: {approval.decision or ''}", previous)
    db.session.commit()
    flash(f"Procurement approval recorded: {decision}.")
    return redirect(url_for("procurement_intake.intake_detail", request_id=request_id))


@procurement_intake_bp.route("/<int:request_id>/order", methods=["POST"])
def create_procurement_order(request_id):
    intake_request = ProcurementIntakeRequest.query.get_or_404(request_id)
    if not ProcurementApproval.query.filter_by(request_id=request_id, status="APPROVED").first():
        flash("An approved procurement decision is required before creating an order.")
        return redirect(url_for("procurement_intake.intake_detail", request_id=request_id)), 409
    order = ProcurementOrder(request=intake_request, order_reference=request.form.get("order_reference", "").strip(), supplier_name=request.form.get("supplier_name", "").strip(), amount=float(request.form.get("amount", 0) or 0), created_by=current_user)
    if not order.order_reference or not order.supplier_name:
        abort(400)
    db.session.add(order)
    intake_request.status = "Delivery Pending"
    record_procurement_event(intake_request, "Procurement order created", "Delivery Pending", f"Order {order.order_reference} created.")
    db.session.commit()
    flash("Procurement order created.")
    return redirect(url_for("procurement_intake.intake_detail", request_id=request_id))


@procurement_intake_bp.route("/<int:request_id>/delivery", methods=["POST"])
def record_procurement_delivery(request_id):
    intake_request = ProcurementIntakeRequest.query.get_or_404(request_id)
    status = request.form.get("status", "DELIVERED")
    if status not in {"PARTIALLY_DELIVERED", "DELIVERED", "DISCREPANCY", "ACCEPTED"}:
        abort(400)
    delivery = ProcurementDelivery(request=intake_request, status=status, quantity=safe_int(request.form.get("quantity"), 0), received_by=current_user, received_at=datetime.utcnow(), comments=request.form.get("comments"))
    db.session.add(delivery)
    intake_request.status = "Completed" if status == "ACCEPTED" else "Delivery Pending"
    record_procurement_event(intake_request, "Delivery recorded", intake_request.status, f"Delivery status: {status}.")
    db.session.commit()
    flash("Delivery record saved.")
    return redirect(url_for("procurement_intake.intake_detail", request_id=request_id))


@procurement_intake_bp.route("/<int:request_id>/close", methods=["POST"])
def close_procurement_request(request_id):
    intake_request = ProcurementIntakeRequest.query.get_or_404(request_id)
    if not ProcurementDelivery.query.filter_by(request_id=request_id, status="ACCEPTED").first():
        flash("Accepted delivery verification is required before closure.")
        return redirect(url_for("procurement_intake.intake_detail", request_id=request_id)), 409
    previous = intake_request.status
    intake_request.status = "Closed"
    add_request_history(intake_request, "RQ closed", from_status=previous, to_status="Closed", comment="Delivery verified and procurement case closed.")
    record_procurement_event(intake_request, "RQ closed", "Closed", "Delivery verified and procurement case closed.", previous)
    db.session.commit()
    flash("RQ formally closed.")
    return redirect(url_for("procurement_intake.intake_detail", request_id=request_id))


internal_control_procurement_bp = Blueprint("internal_control_procurement", __name__, url_prefix="/internal-control/procurement-intake")


@internal_control_procurement_bp.before_request
def protect_internal_control_procurement():
    if not current_user.is_authenticated or not current_user.role_obj or current_user.role_obj.name not in {"System Administrator", "Internal Control Officer", "Internal Audit"}:
        abort(403)


@internal_control_procurement_bp.route("/")
def queue():
    items = ProcurementIntakeRequest.query.filter(ProcurementIntakeRequest.status.in_(["Internal Control Review", "Corrections Required"])).order_by(ProcurementIntakeRequest.updated_at.asc()).all()
    return render_template("procurement_intake_queue.html", items=items, active_filter="Internal Control Review")


@internal_control_procurement_bp.route("/<int:request_id>", methods=["GET", "POST"])
def review(request_id):
    intake_request = ProcurementIntakeRequest.query.get_or_404(request_id)
    if request.method == "POST":
        previous_status = intake_request.status
        new_status = request.form.get("status", "Internal Control Review")
        intake_request.status = new_status
        intake_request.internal_control_notes = request.form.get("internal_control_notes")
        add_request_history(intake_request, "Internal Control review", from_status=previous_status, to_status=new_status, comment=intake_request.internal_control_notes)
        record_procurement_event(intake_request, "Internal Control review", new_status, intake_request.internal_control_notes, previous_status)
        notify_user(intake_request.requester_user_id, f"Internal Control update for {intake_request.request_number}", f"Your procurement request is now {new_status}.")
        db.session.commit()
        log_audit("Internal Control", "review", "procurement_intake_request", f"Reviewed {intake_request.request_number}: {new_status}.", user=current_user)
        flash("Internal Control review saved.")
    return render_template("procurement_request_detail.html", item=intake_request, employee_view=False, internal_control_view=True, officers=[])


@auth_bp.route("/logout")
@login_required
def logout():
    log_audit("Auth", "logout", "user", f"User {current_user.username} logged out.", user=current_user)
    logout_user()
    return redirect(url_for("auth.login"))


@auth_bp.route("/dev-access")
def dev_access():
    if current_app.config.get("ENV") != "development":
        abort(404)

    user = User.query.filter_by(username="devadmin").first()
    if not user:
        role = Role.query.filter_by(name="System Administrator").first()
        if not role:
            abort(500)
        user = User(username="devadmin", email="devadmin@smartchain.test", full_name="Development Administrator", role_id=role.id)
        db.session.add(user)
        db.session.commit()

    login_user(user)
    user.last_login_at = datetime.utcnow()
    log_audit("Auth", "development_access", "user", "Development administrator entered the environment.", user=user)
    return redirect(url_for("dashboard.index"))


# Organisation
org_bp = Blueprint("org", __name__, url_prefix="/org")


@org_bp.route("/departments")
@login_required
def departments():
    items = Department.query.order_by(Department.id).all()
    return render_template("module_list.html", title="Departments", items=items, columns=["Name", "Code", "Description"])


@org_bp.route("/employees")
@login_required
def employees():
    items = Employee.query.order_by(Employee.id).all()
    return render_template("module_list.html", title="Employee 360", items=items, columns=["Employee Number", "Name", "Position", "Status"])


@org_bp.route("/employees/<int:employee_id>")
@login_required
def employee_detail(employee_id):
    employee = Employee.query.get_or_404(employee_id)
    return render_template(
        "employee_detail.html",
        employee=employee,
        department=employee.department,
        directorate=employee.directorate,
        unit=employee.unit,
        role=employee.role,
    )


@org_bp.route("/structure")
@login_required
def organisation_structure():
    departments = Department.query.order_by(Department.name).all()
    return render_template("organisation_structure.html", departments=departments)


@org_bp.route("/employees/new", methods=["GET", "POST"])
@login_required
def create_employee():
    if request.method == "POST":
        payload = request.form
        department_id = safe_int(payload.get("department_id"))
        directorate_id = safe_int(payload.get("directorate_id"))
        unit_id = safe_int(payload.get("unit_id"))
        role_id = safe_int(payload.get("role_id"))
        if None in (department_id, directorate_id, unit_id, role_id):
            flash("Department, directorate, unit, and role are required.")
            departments = Department.query.all()
            directorates = Directorate.query.order_by(Directorate.name).all()
            units = Unit.query.order_by(Unit.name).all()
            roles = Role.query.order_by(Role.name).all()
            return render_template("employee_form.html", departments=departments, directorates=directorates, units=units, roles=roles), 400
        employee = Employee(
            employee_number=payload.get("employee_number"),
            first_name=payload.get("first_name"),
            last_name=payload.get("last_name"),
            title=payload.get("title"),
            position=payload.get("position"),
            email=payload.get("email"),
            phone=payload.get("phone"),
            department_id=department_id,
            directorate_id=directorate_id,
            unit_id=unit_id,
            role_id=role_id,
            account_status=payload.get("account_status", "active"),
        )
        db.session.add(employee)
        db.session.commit()
        log_audit("Employee 360", "create", "employee", f"Created employee {employee.employee_number}.", user=current_user)
        flash("Employee created successfully.")
        return redirect(url_for("org.employees"))
    departments = Department.query.all()
    directorates = Directorate.query.order_by(Directorate.name).all()
    units = Unit.query.order_by(Unit.name).all()
    roles = Role.query.order_by(Role.name).all()
    return render_template("employee_form.html", departments=departments, directorates=directorates, units=units, roles=roles)


# Dashboard

dashboard_bp = Blueprint("dashboard", __name__)


@dashboard_bp.route("/dashboard")
@login_required
def index():
    counts = {
        "employees": Employee.query.count(),
        "demands": DemandRequest.query.count(),
        "requisitions": Requisition.query.filter_by(status="Open").count(),
        "contracts": Contract.query.filter_by(status="Active").count(),
        "active_procurement": ProcurementRequest.query.filter_by(status="Active").count(),
        "suppliers": Supplier.query.count(),
        "tenders": Tender.query.count(),
        "assets": Asset.query.count(),
        "risks": RiskRegister.query.count(),
        "controls": InternalControl.query.count(),
        "security_events": SecurityEvent.query.count(),
        "pending_approvals": ApprovalRecord.query.filter_by(status="Pending").count(),
        "integration_external_live": IntegrationConnection.query.filter_by(mode="PRODUCTION", status="OPERATIONAL").count(),
        "integration_external_total": IntegrationConnection.query.filter(IntegrationConnection.mode.in_(["SANDBOX", "CONFIGURED", "PRODUCTION"])).count(),
        "integration_transactions": IntegrationTransaction.query.count(),
    }
    notifications = Notification.query.filter_by(user_id=current_user.id).order_by(Notification.created_at.desc()).limit(5).all()
    return render_template("dashboard.html", counts=counts, notifications=notifications)


# Demand Management

demand_bp = Blueprint("demand", __name__, url_prefix="/demand")


@demand_bp.route("/")
@login_required
def index():
    items = DemandRequest.query.order_by(DemandRequest.created_at.desc()).all()
    return render_template("module_list.html", title="Demand Management", items=items, columns=["Title", "Priority", "Status", "Estimated Cost"])


@demand_bp.route("/new", methods=["GET", "POST"])
@login_required
def create_demand():
    if request.method == "POST":
        demand = DemandRequest(
            title=request.form.get("title"),
            description=request.form.get("description"),
            responsible_unit=request.form.get("responsible_unit"),
            estimated_cost=float(request.form.get("estimated_cost", 0) or 0),
            priority=request.form.get("priority", "Medium"),
            timing=request.form.get("timing"),
            justification=request.form.get("justification"),
            status="Submitted",
        )
        db.session.add(demand)
        db.session.commit()
        log_audit("Demand Management", "create", "demand_request", f"Created demand request {demand.title}.", user=current_user)
        flash("Demand request created successfully.")
        return redirect(url_for("demand.index"))
    return render_template("demand_form.html")


# APP management
app_bp = Blueprint("app", __name__, url_prefix="/app")


@app_bp.route("/")
@login_required
def index():
    items = AnnualProcurementPlan.query.order_by(AnnualProcurementPlan.created_at.desc()).all()
    return render_template("module_list.html", title="Annual Procurement Plan", items=items, columns=["Item", "Estimated Value", "Method", "Status"])


@app_bp.route("/new", methods=["GET", "POST"])
@login_required
def create_app_item():
    if request.method == "POST":
        item = AnnualProcurementPlan(
            item_name=request.form.get("item_name"),
            estimated_value=float(request.form.get("estimated_value", 0) or 0),
            procurement_method=request.form.get("procurement_method", "Open Tender"),
            responsible_unit=request.form.get("responsible_unit"),
            planned_date=request.form.get("planned_date"),
            budget_allocation=float(request.form.get("budget_allocation", 0) or 0),
            status=request.form.get("status", "Planned"),
        )
        db.session.add(item)
        db.session.commit()
        log_audit("APP", "create", "annual_procurement_plan", f"Created APP item {item.item_name}.", user=current_user)
        flash("APP item created successfully.")
        return redirect(url_for("app.index"))
    return render_template("app_form.html")


# Requisition and procurement
requisition_bp = Blueprint("requisition", __name__, url_prefix="/requisition")


@requisition_bp.route("/")
@login_required
def index():
    status = request.args.get("status")
    query = Requisition.query
    if status:
        query = query.filter_by(status=status)
    items = query.order_by(Requisition.created_at.desc()).all()
    return render_template("module_list.html", title="Requisitions", items=items, columns=["RQ Number", "Requester", "Estimated Value", "Status"], active_filter=status, empty_message=f"No {status.lower()} requisitions" if status else "No requisitions here yet")


@requisition_bp.route("/new", methods=["GET", "POST"])
@login_required
def create_requisition():
    if request.method == "POST":
        requisition = Requisition(
            requisition_number=request.form.get("requisition_number"),
            requester=request.form.get("requester"),
            department=request.form.get("department"),
            directorate=request.form.get("directorate"),
            unit=request.form.get("unit"),
            description=request.form.get("description"),
            motivation=request.form.get("motivation"),
            estimated_value=float(request.form.get("estimated_value", 0) or 0),
            budget=request.form.get("budget"),
            priority=request.form.get("priority", "Medium"),
            status=request.form.get("status", "Submitted"),
        )
        db.session.add(requisition)
        db.session.commit()
        log_audit("Requisition", "create", "requisition", f"Created requisition {requisition.requisition_number}.", user=current_user)
        flash("Requisition created successfully.")
        return redirect(url_for("requisition.index"))
    return render_template("requisition_form.html")


procurement_bp = Blueprint("procurement", __name__, url_prefix="/procurement")


@procurement_bp.route("/")
@login_required
def index():
    status = request.args.get("status")
    query = ProcurementRequest.query
    if status:
        query = query.filter_by(status=status)
    items = query.order_by(ProcurementRequest.created_at.desc()).all()
    return render_template("module_list.html", title="Procurement", items=items, columns=["Title", "Method", "Estimated Value", "Status"], active_filter=status, empty_message=f"No {status.lower()} procurement activities" if status else "No procurement records here yet")


@procurement_bp.route("/new", methods=["GET", "POST"])
@login_required
def create_procurement():
    if request.method == "POST":
        procurement = ProcurementRequest(
            title=request.form.get("title"),
            procurement_method=request.form.get("procurement_method", "Open Tender"),
            estimated_value=float(request.form.get("estimated_value", 0) or 0),
            status=request.form.get("status", "Draft"),
            recommendation=request.form.get("recommendation"),
        )
        db.session.add(procurement)
        db.session.commit()
        log_audit("Procurement", "create", "procurement_request", f"Created procurement request {procurement.title}.", user=current_user)
        flash("Procurement request created successfully.")
        return redirect(url_for("procurement.index"))
    return render_template("procurement_form.html")


# Supplier and comparative
supplier_bp = Blueprint("supplier", __name__, url_prefix="/supplier")


@supplier_bp.route("/")
@login_required
def index():
    items = Supplier.query.order_by(Supplier.created_at.desc()).all()
    return render_template("module_list.html", title="Suppliers", items=items, columns=["Name", "Registration", "Categories", "Status"])


@supplier_bp.route("/new", methods=["GET", "POST"])
@login_required
def create_supplier():
    if request.method == "POST":
        supplier = Supplier(
            name=request.form.get("name"),
            registration_number=request.form.get("registration_number"),
            service_categories=request.form.get("service_categories"),
            status=request.form.get("status", "Active"),
            contact_email=request.form.get("contact_email"),
            phone=request.form.get("phone"),
        )
        db.session.add(supplier)
        db.session.commit()
        log_audit("Supplier Management", "create", "supplier", f"Created supplier {supplier.name}.", user=current_user)
        flash("Supplier created successfully.")
        return redirect(url_for("supplier.index"))
    return render_template("supplier_form.html")


# Tenders and committees

tender_bp = Blueprint("tender", __name__, url_prefix="/tenders")

TENDER_ADMIN_ROLES = {"System Administrator", "SCM Manager", "SCM Officer / Practitioner", "SCM Supervisor"}
TENDER_EVALUATOR_ROLES = TENDER_ADMIN_ROLES | {"BEC Member"}
TENDER_VIEWER_ROLES = TENDER_ADMIN_ROLES | {"Read-Only / Management Viewer"}


def require_tender_admin():
    if not current_user.role_obj or current_user.role_obj.name not in TENDER_ADMIN_ROLES:
        abort(403)


def require_tender_evaluator(tender):
    role_name = current_user.role_obj.name if current_user.role_obj else None
    if role_name in TENDER_ADMIN_ROLES:
        return
    assignment = TenderEvaluatorAssignment.query.filter_by(
        tender_id=tender.id, user_id=current_user.id, active=True
    ).first()
    if role_name not in TENDER_EVALUATOR_ROLES or not assignment:
        abort(403)


def require_tender_workspace_access(tender):
    role_name = current_user.role_obj.name if current_user.role_obj else None
    if role_name in TENDER_VIEWER_ROLES:
        return
    if role_name == "BEC Member" and TenderEvaluatorAssignment.query.filter_by(
        tender_id=tender.id, user_id=current_user.id, active=True
    ).first():
        return
    abort(403)


def record_tender_event(tender, stage, status, details, user=None, document_reference=None):
    event = TenderStageEvent(
        tender=tender,
        stage=stage,
        status=status,
        details=details,
        document_reference=document_reference,
        user=user,
    )
    db.session.add(event)
    return event


def tender_document_folder(tender_id):
    root = current_app.config.get("TENDER_DOCUMENT_FOLDER") or os.path.join(current_app.instance_path, "tender_documents")
    return os.path.join(root, str(tender_id))


def derive_tender_stage(tender):
    if Contract.query.filter(Contract.tender.in_([tender.tender_number, tender.title])).first():
        return "PERFORMANCE"
    if any(outcome.status == "Approved" for outcome in tender.outcomes):
        return "CONTRACT"
    if tender.outcomes:
        return "OUTCOME"
    if any(meeting.committee_type == "BAC" for meeting in tender.committee_meetings) or tender.bac_records:
        return "BAC"
    if any(meeting.committee_type == "BEC" for meeting in tender.committee_meetings) or tender.bec_records:
        return "BEC"
    if any(score.status == "Submitted" for score in tender.evaluations):
        return "EVALUATION"
    if any(requirement.results for requirement in tender.screening_requirements):
        return "SCREENING"
    if tender.bidders:
        return "BID RECEIPT"
    if tender.status.lower() in {"open", "advertised", "published"}:
        return "ADVERTISEMENT"
    return "PLANNING"


@tender_bp.route("/")
@login_required
def index():
    role_name = current_user.role_obj.name if current_user.role_obj else None
    if role_name in TENDER_ADMIN_ROLES:
        items = Tender.query.order_by(Tender.created_at.desc()).all()
    elif role_name == "BEC Member":
        assigned_ids = db.session.query(TenderEvaluatorAssignment.tender_id).filter_by(user_id=current_user.id, active=True)
        items = Tender.query.filter(Tender.id.in_(assigned_ids)).order_by(Tender.created_at.desc()).all()
    elif role_name == "Read-Only / Management Viewer":
        items = Tender.query.order_by(Tender.created_at.desc()).all()
    else:
        abort(403)
    tender_stages = {tender.id: derive_tender_stage(tender) for tender in items}
    return render_template("tender_index.html", items=items, tender_stages=tender_stages)


@tender_bp.route("/new", methods=["GET", "POST"])
@login_required
def create_tender():
    require_tender_admin()
    if request.method == "POST":
        tender = Tender(
            tender_number=request.form.get("tender_number"),
            title=request.form.get("title"),
            description=request.form.get("description"),
            estimated_value=float(request.form.get("estimated_value", 0) or 0),
            closing_date=request.form.get("closing_date"),
            status=request.form.get("status", "Open"),
        )
        tender.profile = TenderProfile(
            directorate=request.form.get("directorate", "").strip() or None,
            responsible_official=request.form.get("responsible_official", "").strip() or None,
            procurement_category=request.form.get("procurement_category", "").strip() or None,
            procurement_type=request.form.get("procurement_type", "").strip() or None,
            opening_date=request.form.get("opening_date", "").strip() or None,
            evaluation_methodology=request.form.get("evaluation_methodology", "").strip() or None,
            mandatory_requirements=request.form.get("mandatory_requirements", "").strip() or None,
            functionality_required=request.form.get("functionality_required") == "true",
            price_preference_method=request.form.get("price_preference_method", "").strip() or None,
            minimum_functionality_threshold=float(request.form.get("minimum_functionality_threshold") or 0) or None,
        )
        db.session.add(tender)
        db.session.commit()
        record_tender_event(tender, "PLANNING", "Created", f"Tender {tender.tender_number} created.", current_user)
        db.session.commit()
        log_audit("Tender Management", "create", "tender", f"Created tender {tender.tender_number}.", user=current_user)
        flash("Tender created successfully.")
        return redirect(url_for("tender.index"))
    return render_template("tender_form.html")


@tender_bp.route("/<int:tender_id>/edit", methods=["GET", "POST"])
@login_required
def edit_tender(tender_id):
    require_tender_admin()
    tender = Tender.query.get_or_404(tender_id)
    if not tender.profile:
        tender.profile = TenderProfile()
        db.session.flush()
    if request.method == "POST":
        tender.title = request.form.get("title", "").strip()
        tender.description = request.form.get("description", "").strip() or None
        tender.estimated_value = float(request.form.get("estimated_value", 0) or 0)
        tender.closing_date = request.form.get("closing_date", "").strip() or None
        tender.status = request.form.get("status", tender.status)
        profile = tender.profile
        profile.directorate = request.form.get("directorate", "").strip() or None
        profile.responsible_official = request.form.get("responsible_official", "").strip() or None
        profile.procurement_category = request.form.get("procurement_category", "").strip() or None
        profile.procurement_type = request.form.get("procurement_type", "").strip() or None
        profile.opening_date = request.form.get("opening_date", "").strip() or None
        profile.evaluation_methodology = request.form.get("evaluation_methodology", "").strip() or None
        profile.mandatory_requirements = request.form.get("mandatory_requirements", "").strip() or None
        profile.functionality_required = request.form.get("functionality_required") == "true"
        profile.price_preference_method = request.form.get("price_preference_method", "").strip() or None
        profile.minimum_functionality_threshold = float(request.form.get("minimum_functionality_threshold") or 0) or None
        db.session.commit()
        record_tender_event(tender, derive_tender_stage(tender), "Updated", "Tender overview and procurement details updated.", current_user)
        db.session.commit()
        log_audit("Tender Management", "update", "tender", f"Updated tender {tender.tender_number}.", user=current_user)
        flash("Tender details updated successfully.")
        return redirect(url_for("tender.workspace", tender_id=tender.id))
    return render_template("tender_form.html", tender=tender, profile=tender.profile)


def calculate_tender_results(tender):
    total_weight = sum(criterion.weight for criterion in tender.criteria)
    results = []
    expected_evaluator_ids = {
        assignment.user_id for assignment in tender.evaluator_assignments if assignment.active
    }
    if not expected_evaluator_ids:
        expected_evaluator_ids = {score.evaluator_id for score in tender.evaluations}
    for bidder in tender.bidders:
        criterion_results = []
        weighted_total = 0.0
        functionality_total = 0.0
        complete = True
        qualified = True
        for criterion in tender.criteria:
            evaluator_scores = [
                score for score in bidder.evaluations if score.criterion_id == criterion.id
            ]
            submitted_by_evaluator = {
                score.evaluator_id: score for score in evaluator_scores
                if score.status == "Submitted" and score.score is not None
            }
            submitted_scores = [score.score for score in submitted_by_evaluator.values()]
            average = sum(submitted_scores) / len(submitted_scores) if submitted_scores else None
            criterion_complete = bool(expected_evaluator_ids) and expected_evaluator_ids.issubset(submitted_by_evaluator)
            if average is None or not criterion_complete:
                complete = False
            elif criterion.max_score:
                weighted_contribution = (average / criterion.max_score) * criterion.weight
                weighted_total += weighted_contribution
                if criterion.config and "function" in (criterion.config.category or "").lower():
                    functionality_total += weighted_contribution
            minimum_score = criterion.config.minimum_qualifying_score if criterion.config else None
            criterion_qualified = average is None or minimum_score is None or average >= minimum_score
            if criterion.pass_fail and average is not None:
                criterion_qualified = average >= criterion.max_score
            if not criterion_qualified and (minimum_score is not None or criterion.pass_fail or criterion.config and criterion.config.mandatory):
                qualified = False
            criterion_results.append({
                "criterion": criterion,
                "average": average,
                "scores": submitted_scores,
                "evaluator_scores": evaluator_scores,
                "complete": criterion_complete,
                "qualified": criterion_qualified,
            })
        profile = tender.profile
        if profile and profile.functionality_required and profile.minimum_functionality_threshold is not None and functionality_total < profile.minimum_functionality_threshold:
            qualified = False
        mandatory_screening = [requirement for requirement in tender.screening_requirements if requirement.mandatory]
        screening_values = []
        for requirement in mandatory_screening:
            screen_result = TenderScreeningResult.query.filter_by(requirement_id=requirement.id, bidder_id=bidder.id).first()
            screening_values.append(screen_result.result if screen_result else "PENDING")
            if screen_result and screen_result.result == "FAIL":
                qualified = False
        results.append({
            "bidder": bidder,
            "criteria": criterion_results,
            "total": weighted_total if total_weight else None,
            "complete": complete and bool(tender.criteria),
            "qualified": qualified,
            "functionality_total": functionality_total,
            "screening_status": "FAIL" if "FAIL" in screening_values else "PASS" if screening_values and all(value == "PASS" for value in screening_values) else "PENDING" if screening_values else "NOT CONFIGURED",
        })
    complete_results = [result for result in results if result["complete"] and result["qualified"]]
    previous_total = None
    previous_rank = None
    for position, result in enumerate(sorted(complete_results, key=lambda item: item["total"], reverse=True), start=1):
        if previous_total is None or result["total"] != previous_total:
            previous_rank = position
        result["rank"] = previous_rank
        previous_total = result["total"]
    return results


@tender_bp.route("/<int:tender_id>")
@login_required
def workspace(tender_id):
    tender = Tender.query.get_or_404(tender_id)
    require_tender_workspace_access(tender)
    users = User.query.filter(User.is_active.is_(True)).order_by(User.full_name).all()
    assignments = TenderEvaluatorAssignment.query.filter_by(tender_id=tender.id, active=True).all()
    tender_contracts = Contract.query.filter_by(tender=tender.tender_number).order_by(Contract.created_at.desc()).all()
    ordered_criteria = sorted(tender.criteria, key=lambda item: item.config.position if item.config else item.id)
    mandatory_requirements = [requirement for requirement in tender.screening_requirements if requirement.mandatory]
    bidder_screening = {
        bidder.id: [
            TenderScreeningResult.query.filter_by(requirement_id=requirement.id, bidder_id=bidder.id).first()
            for requirement in mandatory_requirements
        ]
        for bidder in tender.bidders
    }
    passed_screening = sum(
        1 for results_for_bidder in bidder_screening.values()
        if results_for_bidder and all(result and result.result == "PASS" for result in results_for_bidder)
    )
    expected_evaluations = len(tender.bidders) * len(tender.criteria) * len(assignments)
    submitted_evaluations = EvaluatorScore.query.filter_by(tender_id=tender.id, status="Submitted").count()
    latest_bec = next((meeting for meeting in reversed(tender.committee_meetings) if meeting.committee_type == "BEC"), None)
    latest_bac = next((meeting for meeting in reversed(tender.committee_meetings) if meeting.committee_type == "BAC"), None)
    audit_entries = AuditLog.query.filter(AuditLog.details.contains(tender.tender_number)).order_by(AuditLog.created_at.desc()).limit(50).all()
    dashboard = {
        "bidder_count": len(tender.bidders),
        "screened_pass": passed_screening,
        "screened_pending": max(0, len(tender.bidders) - passed_screening),
        "evaluation_submitted": submitted_evaluations,
        "evaluation_expected": expected_evaluations,
        "bec_status": latest_bec.status if latest_bec else "Not scheduled",
        "bac_status": latest_bac.status if latest_bac else "Pending",
        "document_count": len(tender.documents),
        "audit_count": len(audit_entries),
        "bec_reports": len([report for report in tender.generated_reports if report.committee_type == "BEC"]),
        "bac_reports": len([report for report in tender.generated_reports if report.committee_type == "BAC"]),
    }
    return render_template(
        "tender_workspace.html",
        tender=tender,
        results=calculate_tender_results(tender),
        users=users,
        assignments=assignments,
        ordered_criteria=ordered_criteria,
        stage=derive_tender_stage(tender),
        dashboard=dashboard,
        audit_entries=audit_entries,
        tender_contracts=tender_contracts,
    )


@tender_bp.route("/<int:tender_id>/documents", methods=["GET", "POST"])
@login_required
def tender_documents(tender_id):
    tender = Tender.query.get_or_404(tender_id)
    require_tender_workspace_access(tender)
    allowed_extensions = {"pdf", "doc", "docx", "xls", "xlsx", "csv"}
    if request.method == "POST":
        require_tender_admin()
        upload = request.files.get("document")
        if not upload or not upload.filename:
            flash("Choose a tender document to upload.")
            return redirect(url_for("tender.tender_documents", tender_id=tender.id))
        original_filename = secure_filename(upload.filename)
        extension = original_filename.rsplit(".", 1)[-1].lower() if "." in original_filename else ""
        if extension not in allowed_extensions:
            flash("Tender documents must be PDF, Word, Excel, or CSV files.")
            return redirect(url_for("tender.tender_documents", tender_id=tender.id)), 400
        category = request.form.get("category", "Supporting document").strip() or "Supporting document"
        prior_versions = TenderDocument.query.filter_by(
            tender_id=tender.id, category=category, original_filename=original_filename
        ).all()
        version = max((document.version for document in prior_versions), default=0) + 1
        stored_filename = f"{tender.id}_{uuid.uuid4().hex}_{original_filename}"
        upload_folder = tender_document_folder(tender.id)
        os.makedirs(upload_folder, exist_ok=True)
        upload.save(os.path.join(upload_folder, stored_filename))
        document = TenderDocument(
            tender=tender,
            category=category,
            original_filename=original_filename,
            stored_filename=stored_filename,
            version=version,
            description=request.form.get("description", "").strip() or None,
            uploaded_by=current_user,
        )
        db.session.add(document)
        record_tender_event(tender, "TENDER DOCUMENTS", f"Version {version} uploaded", f"{category}: {original_filename}", current_user, stored_filename)
        db.session.commit()
        log_audit("Tender Documents", "upload", "tender_document", f"Uploaded {original_filename} version {version} for {tender.tender_number}.", user=current_user)
        flash(f"Document uploaded as version {version}; earlier versions were preserved.")
        return redirect(url_for("tender.tender_documents", tender_id=tender.id))
    documents = TenderDocument.query.filter_by(tender_id=tender.id).order_by(TenderDocument.category, TenderDocument.original_filename, TenderDocument.version.desc()).all()
    return render_template("tender_documents.html", tender=tender, documents=documents)


@tender_bp.route("/<int:tender_id>/documents/<int:document_id>/download")
@login_required
def download_tender_document(tender_id, document_id):
    tender = Tender.query.get_or_404(tender_id)
    require_tender_workspace_access(tender)
    document = TenderDocument.query.filter_by(id=document_id, tender_id=tender.id).first_or_404()
    upload_folder = tender_document_folder(tender.id)
    return send_from_directory(upload_folder, document.stored_filename, as_attachment=True, download_name=document.original_filename)


@tender_bp.route("/<int:tender_id>/committees/<committee_type>/meetings", methods=["GET", "POST"])
@login_required
def tender_committee_meetings(tender_id, committee_type):
    if committee_type not in {"BEC", "BAC"}:
        abort(404)
    tender = Tender.query.get_or_404(tender_id)
    if request.method == "POST":
        require_tender_admin()
    else:
        require_tender_workspace_access(tender)
        if current_user.role_obj and current_user.role_obj.name == "BEC Member" and committee_type != "BEC":
            abort(403)
    if request.method == "POST":
        meeting_at_value = request.form.get("meeting_at", "").strip()
        try:
            meeting_at = datetime.fromisoformat(meeting_at_value) if meeting_at_value else None
        except ValueError:
            flash("Enter a valid meeting date and time.")
            return redirect(url_for("tender.tender_committee_meetings", tender_id=tender.id, committee_type=committee_type)), 400
        meeting = TenderCommitteeMeeting(
            tender=tender,
            committee_type=committee_type,
            meeting_number=request.form.get("meeting_number", "").strip() or None,
            meeting_at=meeting_at,
            venue=request.form.get("venue", "").strip() or None,
            chair=request.form.get("chair", "").strip() or None,
            secretariat=request.form.get("secretariat", "").strip() or None,
            agenda_reference=request.form.get("agenda_reference", "").strip() or None,
            minutes=request.form.get("minutes", "").strip() or None,
            resolution=request.form.get("resolution", "").strip() or None,
            status=request.form.get("status", "Planned"),
            created_by=current_user,
        )
        db.session.add(meeting)
        db.session.flush()
        member_names = request.form.getlist("member_name")
        member_roles = request.form.getlist("member_role")
        member_organizations = request.form.getlist("member_organization")
        for index, member_name in enumerate(member_names):
            if member_name.strip():
                db.session.add(TenderCommitteeMember(
                    meeting=meeting,
                    name=member_name.strip(),
                    role=member_roles[index].strip() if index < len(member_roles) else None,
                    organization=member_organizations[index].strip() if index < len(member_organizations) else None,
                ))
        record_tender_event(tender, committee_type, status=meeting.status, details=f"{committee_type} meeting {meeting.meeting_number or 'recorded'}.", user=current_user)
        db.session.commit()
        log_audit(committee_type, "meeting", "tender_committee_meeting", f"Recorded {committee_type} meeting for {tender.tender_number}.", user=current_user)
        flash(f"{committee_type} meeting saved in this tender workspace.")
        return redirect(url_for("tender.tender_committee_meetings", tender_id=tender.id, committee_type=committee_type))
    meetings = TenderCommitteeMeeting.query.filter_by(tender_id=tender.id, committee_type=committee_type).order_by(TenderCommitteeMeeting.meeting_at.desc()).all()
    return render_template("tender_committee_meetings.html", tender=tender, committee_type=committee_type, meetings=meetings)


@tender_bp.route("/<int:tender_id>/committees/meetings/<int:meeting_id>/attendance", methods=["POST"])
@login_required
def update_tender_meeting_attendance(tender_id, meeting_id):
    require_tender_admin()
    meeting = TenderCommitteeMeeting.query.filter_by(id=meeting_id, tender_id=tender_id).first_or_404()
    attended_ids = {safe_int(value) for value in request.form.getlist("attended_member_ids")}
    for member in meeting.members:
        member.attended = member.id in attended_ids
    db.session.commit()
    log_audit(meeting.committee_type, "attendance", "tender_committee_meeting", f"Recorded attendance for {meeting.committee_type} meeting {meeting.meeting_number or meeting.id}.", user=current_user)
    flash("Committee attendance recorded.")
    return redirect(url_for("tender.tender_committee_meetings", tender_id=tender_id, committee_type=meeting.committee_type))


@tender_bp.route("/<int:tender_id>/committees/<committee_type>/report.pdf")
@login_required
def generate_tender_committee_report(tender_id, committee_type):
    if committee_type not in {"BEC", "BAC"}:
        abort(404)
    tender = Tender.query.get_or_404(tender_id)
    require_tender_workspace_access(tender)
    if current_user.role_obj and current_user.role_obj.name == "BEC Member" and committee_type != "BEC":
        abort(403)
    meetings = TenderCommitteeMeeting.query.filter_by(tender_id=tender.id, committee_type=committee_type).order_by(TenderCommitteeMeeting.meeting_at).all()
    if not meetings:
        flash(f"Record a {committee_type} meeting before generating its report.")
        return redirect(url_for("tender.tender_committee_meetings", tender_id=tender.id, committee_type=committee_type)), 409
    version = (db.session.query(db.func.max(TenderReportRecord.version)).filter_by(tender_id=tender.id, committee_type=committee_type).scalar() or 0) + 1
    report_reference = f"{tender.tender_number}-{committee_type}-REPORT-v{version}"
    report_record = TenderReportRecord(tender=tender, committee_type=committee_type, version=version, report_reference=report_reference, generated_by=current_user)
    db.session.add(report_record)
    record_tender_event(tender, committee_type, "Report generated", f"{report_reference} generated from current tender records.", current_user, report_reference)
    db.session.commit()
    log_audit(committee_type, "report_generated", "tender_committee_report", f"Generated {report_reference}.", user=current_user)

    from xml.sax.saxutils import escape
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.lib.units import mm
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    output = BytesIO()
    pdf = SimpleDocTemplate(output, pagesize=A4, rightMargin=16 * mm, leftMargin=16 * mm, topMargin=16 * mm, bottomMargin=16 * mm)
    styles = getSampleStyleSheet()
    story = [
        Paragraph(f"{committee_type} Evaluation Report", styles["Title"]),
        Paragraph(f"Report reference: {escape(report_reference)} · Generated {datetime.utcnow().strftime('%Y-%m-%d %H:%M UTC')} · {escape(current_user.full_name)}", styles["Normal"]),
        Spacer(1, 8 * mm),
        Paragraph("Tender information", styles["Heading2"]),
        Paragraph(f"{escape(tender.tender_number)} · {escape(tender.title)}", styles["BodyText"]),
        Paragraph(f"Status: {escape(tender.status)} · Current stage: {escape(derive_tender_stage(tender))} · Estimated value: R {tender.estimated_value or 0:,.2f}", styles["BodyText"]),
        Paragraph(f"Description: {escape(tender.description or 'Not recorded')}", styles["BodyText"]),
        Spacer(1, 5 * mm),
        Paragraph("Evaluation methodology and criteria", styles["Heading2"]),
    ]
    profile = tender.profile
    story.append(Paragraph(f"Methodology: {escape((profile.evaluation_methodology if profile else None) or 'Not recorded')} · Procurement type: {escape((profile.procurement_type if profile else None) or 'Not recorded')}", styles["BodyText"]))
    criterion_rows = [["Criterion", "Weight", "Framework", "Maximum", "Threshold"]]
    for criterion in tender.criteria:
        criterion_rows.append([criterion.name, f"{criterion.weight:g}%", criterion.scoring_framework, str(criterion.max_score), str(criterion.config.minimum_qualifying_score or "-") if criterion.config else "-"])
    story.append(Table(criterion_rows, repeatRows=1, style=TableStyle([("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#dfe8e3")), ("GRID", (0, 0), (-1, -1), 0.4, colors.grey), ("VALIGN", (0, 0), (-1, -1), "TOP")])))
    story.extend([Spacer(1, 5 * mm), Paragraph("Bidder register and consolidated evaluation", styles["Heading2"])])
    result_by_bidder = {result["bidder"].id: result for result in calculate_tender_results(tender)}
    bidder_rows = [["Bidder", "Supplier reference", "Bid price", "Screening"]]
    bidder_rows[0].extend([criterion.name for criterion in tender.criteria])
    bidder_rows[0].extend(["Weighted total", "Evaluation position"])
    for bidder in tender.bidders:
        screening = [record.result for record in bidder.screening_results]
        result = result_by_bidder.get(bidder.id, {})
        position = f"Current position {result['rank']}" if result.get("rank") else "Evaluation in progress"
        criterion_values = [
            f"{criterion_result['average']:.2f}" if criterion_result["average"] is not None else "Pending"
            for criterion_result in result.get("criteria", [])
        ]
        weighted_total = f"{result['total']:.2f}" if result.get("total") is not None else "Pending"
        bidder_rows.append([
            bidder.legal_name,
            bidder.supplier_reference or "-",
            f"R {bidder.bid_price or 0:,.2f}",
            ", ".join(screening) or "Pending",
            *criterion_values,
            weighted_total,
            position,
        ])
    story.append(Table(bidder_rows, repeatRows=1, style=TableStyle([("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#dfe8e3")), ("GRID", (0, 0), (-1, -1), 0.4, colors.grey), ("VALIGN", (0, 0), (-1, -1), "TOP")])))
    for meeting in meetings:
        story.extend([Spacer(1, 5 * mm), Paragraph(f"Meeting {escape(meeting.meeting_number or str(meeting.id))}", styles["Heading2"]), Paragraph(f"Date: {meeting.meeting_at.strftime('%Y-%m-%d %H:%M') if meeting.meeting_at else 'Not recorded'} · Venue: {escape(meeting.venue or 'Not recorded')} · Chair: {escape(meeting.chair or 'Not recorded')} · Secretariat: {escape(meeting.secretariat or 'Not recorded')}", styles["BodyText"]), Paragraph(f"Agenda reference: {escape(meeting.agenda_reference or 'Not recorded')}", styles["BodyText"]), Paragraph("Members and attendance", styles["Heading3"])])
        for member in meeting.members:
            story.append(Paragraph(f"{escape(member.name)} · {escape(member.role or 'Member')} · {escape(member.organization or '')} · {'Present' if member.attended else 'Not recorded as present'}", styles["BodyText"]))
        story.extend([Paragraph("Minutes and deliberations", styles["Heading3"]), Paragraph(escape(meeting.minutes or "Not recorded"), styles["BodyText"]), Paragraph("Recommendation / resolution", styles["Heading3"]), Paragraph(escape(meeting.resolution or "Not recorded"), styles["BodyText"])])
    story.extend([Spacer(1, 5 * mm), Paragraph("Tender document references", styles["Heading2"])])
    for document in tender.documents:
        story.append(Paragraph(f"{escape(document.category)} · {escape(document.original_filename)} · version {document.version}", styles["BodyText"]))
    if tender.outcomes:
        story.extend([Spacer(1, 5 * mm), Paragraph("Recorded outcomes", styles["Heading2"])])
        for outcome in tender.outcomes:
            story.append(Paragraph(f"{escape(outcome.decision_type)} · {escape(outcome.status)} · {outcome.decided_at.strftime('%Y-%m-%d %H:%M')} · {escape(outcome.rationale)}", styles["BodyText"]))
    pdf.build(story)
    return Response(output.getvalue(), mimetype="application/pdf", headers={"Content-Disposition": f"attachment; filename={report_reference}.pdf"})


@tender_bp.route("/<int:tender_id>/committees/<committee_type>/report.xlsx")
@login_required
def export_tender_committee_report_xlsx(tender_id, committee_type):
    if committee_type not in {"BEC", "BAC"}:
        abort(404)
    tender = Tender.query.get_or_404(tender_id)
    require_tender_workspace_access(tender)
    if current_user.role_obj and current_user.role_obj.name == "BEC Member" and committee_type != "BEC":
        abort(403)
    meetings = TenderCommitteeMeeting.query.filter_by(tender_id=tender.id, committee_type=committee_type).order_by(TenderCommitteeMeeting.meeting_at).all()
    if not meetings:
        flash(f"Record a {committee_type} meeting before generating its report.")
        return redirect(url_for("tender.tender_committee_meetings", tender_id=tender.id, committee_type=committee_type)), 409
    version = (db.session.query(db.func.max(TenderReportRecord.version)).filter_by(tender_id=tender.id, committee_type=committee_type).scalar() or 0) + 1
    report_reference = f"{tender.tender_number}-{committee_type}-REPORT-v{version}"
    db.session.add(TenderReportRecord(tender=tender, committee_type=committee_type, version=version, report_reference=report_reference, generated_by=current_user))
    record_tender_event(tender, committee_type, "Spreadsheet report generated", report_reference, current_user, report_reference)
    db.session.commit()
    log_audit(committee_type, "spreadsheet_report_generated", "tender_committee_report", f"Generated {report_reference}.xlsx.", user=current_user)

    from openpyxl import Workbook

    workbook = Workbook()
    overview = workbook.active
    overview.title = "Tender overview"
    overview.append(["Report reference", report_reference])
    overview.append(["Tender number", tender.tender_number])
    overview.append(["Title", tender.title])
    overview.append(["Description", tender.description or ""])
    overview.append(["Status", tender.status])
    overview.append(["Current stage", derive_tender_stage(tender)])
    overview.append(["Generated by", current_user.full_name])
    overview.append(["Generated at UTC", datetime.utcnow().strftime("%Y-%m-%d %H:%M")])

    criteria_sheet = workbook.create_sheet("Criteria")
    criteria_sheet.append(["Criterion", "Category", "Weight", "Framework", "Maximum", "Threshold", "Mandatory"])
    for criterion in tender.criteria:
        criteria_sheet.append([
            criterion.name,
            criterion.config.category if criterion.config else "",
            criterion.weight,
            criterion.scoring_framework,
            criterion.max_score,
            criterion.config.minimum_qualifying_score if criterion.config else None,
            criterion.config.mandatory if criterion.config else False,
        ])

    bidders_sheet = workbook.create_sheet("Bidders and evaluation")
    bidders_sheet.append(["Bidder", "Trading name", "Registration number", "Supplier reference", "Contact email", "Phone", "Submitted at", "Bid price", "Screening", "Criterion", "Evaluator", "Score", "Score status", "Current position"])
    result_by_bidder = {result["bidder"].id: result for result in calculate_tender_results(tender)}
    for bidder in tender.bidders:
        bidder_result = result_by_bidder.get(bidder.id, {})
        score_rows = [
            (criterion_result, score)
            for criterion_result in bidder_result.get("criteria", [])
            for score in criterion_result["evaluator_scores"]
        ] or [(None, None)]
        for criterion_result, score in score_rows:
            bidders_sheet.append([
                bidder.legal_name,
                bidder.profile.trading_name if bidder.profile else "",
                bidder.profile.registration_number if bidder.profile else "",
                bidder.supplier_reference or "",
                bidder.profile.contact_email if bidder.profile else "",
                bidder.profile.phone if bidder.profile else "",
                bidder.submitted_at.isoformat() if bidder.submitted_at else "",
                bidder.bid_price,
                bidder_result.get("screening_status", "PENDING"),
                criterion_result["criterion"].name if criterion_result else "",
                score.evaluator.full_name if score else "",
                score.score if score else None,
                score.status if score else "Not scored",
                f"Current position {bidder_result['rank']}" if bidder_result.get("rank") else "Evaluation in progress",
            ])

    meetings_sheet = workbook.create_sheet(f"{committee_type} meetings")
    meetings_sheet.append(["Meeting", "Date/time", "Venue", "Chair", "Secretariat", "Agenda reference", "Member", "Role", "Organisation", "Attendance", "Minutes", "Resolution"])
    for meeting in meetings:
        meeting_members = meeting.members or [None]
        for member in meeting_members:
            meetings_sheet.append([
                meeting.meeting_number or meeting.id,
                meeting.meeting_at.isoformat() if meeting.meeting_at else "",
                meeting.venue or "",
                meeting.chair or "",
                meeting.secretariat or "",
                meeting.agenda_reference or "",
                member.name if member else "",
                member.role if member else "",
                member.organization if member else "",
                member.attended if member else None,
                meeting.minutes or "",
                meeting.resolution or "",
            ])

    output = BytesIO()
    workbook.save(output)
    return Response(
        output.getvalue(),
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename={report_reference}.xlsx"},
    )


@tender_bp.route("/<int:tender_id>/outcome", methods=["GET", "POST"])
@login_required
def record_tender_outcome(tender_id):
    require_tender_admin()
    tender = Tender.query.get_or_404(tender_id)
    statuses = ["Proposed", "Approved", "Rejected", "Cancelled", "Retender", "No Award"]
    if request.method == "POST":
        status = request.form.get("status", "")
        decision_type = request.form.get("decision_type", "").strip()
        rationale = request.form.get("rationale", "").strip()
        bidder = TenderBidder.query.filter_by(id=safe_int(request.form.get("bidder_id")), tender_id=tender.id).first() if request.form.get("bidder_id") else None
        if status not in statuses or not decision_type or not rationale or (status == "Approved" and not bidder):
            flash("Select a valid decision, give a rationale, and select a bidder for an approved outcome.")
            return render_template("tender_outcome.html", tender=tender, statuses=statuses), 400
        outcome = TenderOutcome(tender=tender, status=status, decision_type=decision_type, bidder=bidder, rationale=rationale, decided_by=current_user)
        db.session.add(outcome)
        record_tender_event(tender, "OUTCOME", status, f"{decision_type}: {rationale}", current_user)
        db.session.commit()
        log_audit("Tender Outcome", "record", "tender_outcome", f"Recorded {status} outcome for {tender.tender_number}.", user=current_user)
        flash("Authorised tender outcome recorded.")
        return redirect(url_for("tender.workspace", tender_id=tender.id))
    return render_template("tender_outcome.html", tender=tender, statuses=statuses)


@tender_bp.route("/<int:tender_id>/contract", methods=["GET", "POST"])
@login_required
def create_tender_contract(tender_id):
    require_tender_admin()
    tender = Tender.query.get_or_404(tender_id)
    approved_outcome = next((outcome for outcome in reversed(tender.outcomes) if outcome.status == "Approved" and outcome.bidder), None)
    if not approved_outcome:
        flash("Record an approved, bidder-specific outcome before creating the tender contract.")
        return redirect(url_for("tender.record_tender_outcome", tender_id=tender.id)), 409
    if request.method == "POST":
        contract = Contract(
            contract_number=request.form.get("contract_number", "").strip(),
            supplier=approved_outcome.bidder.legal_name,
            tender=tender.tender_number,
            start_date=request.form.get("start_date", "").strip() or None,
            end_date=request.form.get("end_date", "").strip() or None,
            value=float(request.form.get("value") or approved_outcome.bidder.bid_price or 0),
            status=request.form.get("status", "Draft"),
        )
        db.session.add(contract)
        record_tender_event(tender, "CONTRACT", contract.status, f"Contract {contract.contract_number} created for {contract.supplier}.", current_user)
        db.session.commit()
        log_audit("Contracts", "tender_handoff", "contract", f"Created contract {contract.contract_number} from tender {tender.tender_number}.", user=current_user)
        flash("Tender contract created.")
        return redirect(url_for("tender.workspace", tender_id=tender.id))
    return render_template("tender_contract_form.html", tender=tender, outcome=approved_outcome)


@tender_bp.route("/<int:tender_id>/bidders/new", methods=["GET", "POST"])
@login_required
def create_tender_bidder(tender_id):
    require_tender_admin()
    tender = Tender.query.get_or_404(tender_id)
    if request.method == "POST":
        bidder = TenderBidder(
            tender=tender,
            legal_name=request.form.get("legal_name"),
            supplier_reference=request.form.get("supplier_reference"),
            registration_details=request.form.get("registration_details"),
            bid_price=float(request.form.get("bid_price", 0) or 0),
            compliance_status=request.form.get("compliance_status", "Pending"),
            comments=request.form.get("comments"),
            submitted_at=datetime.fromisoformat(request.form["submitted_at"]) if request.form.get("submitted_at") else None,
        )
        bidder.profile = TenderBidderProfile(
            trading_name=request.form.get("trading_name", "").strip() or None,
            registration_number=request.form.get("registration_number", "").strip() or None,
            contact_name=request.form.get("contact_name", "").strip() or None,
            contact_email=request.form.get("contact_email", "").strip() or None,
            phone=request.form.get("phone", "").strip() or None,
        )
        db.session.add(bidder)
        db.session.commit()
        record_tender_event(tender, "BID RECEIPT", "Bidder registered", f"Registered bidder {bidder.legal_name}.", current_user)
        db.session.commit()
        log_audit("Tender Management", "create", "tender_bidder", f"Registered bidder {bidder.legal_name} for {tender.tender_number}.", user=current_user)
        flash("Bidder registered successfully.")
        return redirect(url_for("tender.workspace", tender_id=tender.id))
    return render_template("tender_bidder_form.html", tender=tender)


@tender_bp.route("/<int:tender_id>/bidders/import", methods=["GET", "POST"])
@login_required
def import_tender_bidders(tender_id):
    require_tender_admin()
    tender = Tender.query.get_or_404(tender_id)
    if request.method == "GET":
        return render_template("tender_bidder_import.html", tender=tender)

    upload = request.files.get("bidder_file")
    if not upload or not upload.filename:
        flash("Choose a bidder file to import.")
        return redirect(url_for("tender.import_tender_bidders", tender_id=tender.id))
    filename = secure_filename(upload.filename)
    extension = filename.rsplit(".", 1)[-1].lower() if "." in filename else ""
    if extension not in {"csv", "xlsx", "pdf", "docx"}:
        flash("Supported bidder import formats are CSV, XLSX, PDF and DOCX.")
        return redirect(url_for("tender.import_tender_bidders", tender_id=tender.id)), 400

    rows = []
    confidence = "High"
    try:
        if extension == "csv":
            content = upload.stream.read().decode("utf-8-sig")
            reader = csv.DictReader(StringIO(content))
            if not reader.fieldnames:
                raise ValueError("The CSV must contain a header row.")
            normalized_headers = {header.strip().lower().replace(" ", "_") for header in reader.fieldnames if header}
            has_bidder_columns = bool(normalized_headers & {"legal_name", "company_name", "bidder", "company", "name"})
            for row in reader:
                row = {key.strip().lower().replace(" ", "_"): (value or "").strip() for key, value in row.items() if key}
                rows.append({
                    "legal_name": next((row.get(key) for key in ("legal_name", "company_name", "bidder", "company", "name") if row.get(key)), ""),
                    "supplier_reference": next((row.get(key) for key in ("supplier_reference", "registration_number", "bid_reference", "reference") if row.get(key)), ""),
                    "bid_price": next((row.get(key) for key in ("bid_price", "price", "amount") if row.get(key)), ""),
                    "trading_name": row.get("trading_name", ""),
                    "registration_number": row.get("registration_number", ""),
                    "contact_name": row.get("contact_name", ""),
                    "contact_email": row.get("contact_email", ""),
                    "phone": row.get("phone", ""),
                    "submitted_at": row.get("submitted_at", ""),
                })
            if not has_bidder_columns:
                raise ValueError("CSV needs a company_name, legal_name, bidder, company or name column.")
        elif extension == "xlsx":
            from openpyxl import load_workbook

            workbook = load_workbook(upload.stream, read_only=True, data_only=True)
            sheet = workbook.active
            values = sheet.iter_rows(values_only=True)
            headers = [str(value or "").strip().lower().replace(" ", "_") for value in next(values, ())]
            if not headers:
                raise ValueError("The spreadsheet must contain a header row.")
            for values_row in values:
                row = {header: str(value or "").strip() for header, value in zip(headers, values_row) if header}
                rows.append({
                    "legal_name": next((row.get(key) for key in ("legal_name", "company_name", "bidder", "company", "name") if row.get(key)), ""),
                    "supplier_reference": next((row.get(key) for key in ("supplier_reference", "registration_number", "bid_reference", "reference") if row.get(key)), ""),
                    "bid_price": next((row.get(key) for key in ("bid_price", "price", "amount") if row.get(key)), ""),
                    "trading_name": row.get("trading_name", ""),
                    "registration_number": row.get("registration_number", ""),
                    "contact_name": row.get("contact_name", ""),
                    "contact_email": row.get("contact_email", ""),
                    "phone": row.get("phone", ""),
                    "submitted_at": row.get("submitted_at", ""),
                })
        else:
            confidence = "Low · confirm each detected line"
            if extension == "pdf":
                from pypdf import PdfReader

                text = "\n".join(page.extract_text() or "" for page in PdfReader(upload.stream).pages)
            else:
                from docx import Document as DocxDocument

                text = "\n".join(paragraph.text for paragraph in DocxDocument(upload.stream).paragraphs)
            for line in text.splitlines():
                cells = [cell.strip() for cell in line.split("|") if cell.strip()] if "|" in line else [cell.strip() for cell in line.split(",")]
                if not cells or not cells[0]:
                    continue
                rows.append({"legal_name": cells[0], "supplier_reference": cells[1] if len(cells) > 1 else "", "bid_price": cells[2] if len(cells) > 2 else "", "trading_name": "", "registration_number": "", "contact_name": "", "contact_email": "", "phone": "", "submitted_at": ""})
    except (UnicodeDecodeError, ValueError, OSError) as error:
        flash(f"Could not read bidder file: {error}")
        return redirect(url_for("tender.import_tender_bidders", tender_id=tender.id)), 400

    rows = [row for row in rows if any(value.strip() for value in row.values())]
    if not rows:
        flash("No bidder rows could be read. Scanned PDFs and images are not OCR processed.")
        return redirect(url_for("tender.import_tender_bidders", tender_id=tender.id)), 400
    session["tender_bidder_import"] = {"tender_id": tender.id, "user_id": current_user.id, "rows": rows}
    return render_template("tender_bidder_import_review.html", tender=tender, rows=rows, confidence=confidence)


@tender_bp.route("/<int:tender_id>/bidders/import/confirm", methods=["POST"])
@login_required
def confirm_tender_bidder_import(tender_id):
    require_tender_admin()
    tender = Tender.query.get_or_404(tender_id)
    staged = session.get("tender_bidder_import")
    if not staged or staged.get("tender_id") != tender.id or staged.get("user_id") != current_user.id:
        flash("The bidder import review has expired. Upload the file again.")
        return redirect(url_for("tender.import_tender_bidders", tender_id=tender.id)), 400
    names = request.form.getlist("legal_name")
    references = request.form.getlist("supplier_reference")
    prices = request.form.getlist("bid_price")
    profile_fields = {
        field: request.form.getlist(field)
        for field in ("trading_name", "registration_number", "contact_name", "contact_email", "phone", "submitted_at")
    }
    selected = {safe_int(value) for value in request.form.getlist("include_row")}
    bidders = []
    errors = []
    for index, row in enumerate(staged["rows"]):
        if index not in selected:
            continue
        legal_name = names[index].strip() if index < len(names) else ""
        reference = references[index].strip() if index < len(references) else ""
        price_text = prices[index].strip() if index < len(prices) else ""
        if not legal_name:
            errors.append(f"Row {index + 1}: company name is required")
            continue
        try:
            bid_price = float(price_text or 0)
            if not math.isfinite(bid_price) or bid_price < 0:
                raise ValueError
        except ValueError:
            errors.append(f"Row {index + 1}: bid price must be a non-negative number")
            continue
        bidder = TenderBidder(tender=tender, legal_name=legal_name, supplier_reference=reference or None, bid_price=bid_price)
        profile_values = {
            field: (profile_fields[field][index].strip() if index < len(profile_fields[field]) else row.get(field, "").strip())
            for field in profile_fields
        }
        submitted_at = None
        if profile_values["submitted_at"]:
            try:
                submitted_at = datetime.fromisoformat(profile_values["submitted_at"])
            except ValueError:
                errors.append(f"Row {index + 1}: submission date/time must use ISO format.")
                continue
        bidder.submitted_at = submitted_at
        bidder.profile = TenderBidderProfile(
            trading_name=profile_values["trading_name"] or None,
            registration_number=profile_values["registration_number"] or None,
            contact_name=profile_values["contact_name"] or None,
            contact_email=profile_values["contact_email"] or None,
            phone=profile_values["phone"] or None,
        )
        bidders.append(bidder)
    if errors or not bidders:
        flash("Nothing imported. " + (" | ".join(errors[:5]) if errors else "Select at least one valid bidder."))
        return redirect(url_for("tender.import_tender_bidders", tender_id=tender.id)), 400
    db.session.add_all(bidders)
    record_tender_event(tender, "BID RECEIPT", "Imported", f"{len(bidders)} bidder records imported after user review.", current_user)
    db.session.commit()
    log_audit("Tender Management", "bulk_import", "tender_bidder", f"Imported {len(bidders)} bidders to {tender.tender_number} after review.", user=current_user)
    session.pop("tender_bidder_import", None)
    flash(f"Imported {len(bidders)} reviewed bidder records.")
    return redirect(url_for("tender.workspace", tender_id=tender.id))


@tender_bp.route("/<int:tender_id>/bidders/<int:bidder_id>/submissions/new", methods=["GET", "POST"])
@login_required
def create_bid_submission(tender_id, bidder_id):
    require_tender_admin()
    tender = Tender.query.get_or_404(tender_id)
    bidder = TenderBidder.query.filter_by(id=bidder_id, tender_id=tender.id).first_or_404()
    if request.method == "POST":
        document_category = request.form.get("document_category", "Bid Submission")
        upload = request.files.get("document_file")
        document_reference = request.form.get("document_reference", "").strip() or None
        if upload and upload.filename:
            original_filename = secure_filename(upload.filename)
            extension = original_filename.rsplit(".", 1)[-1].lower() if "." in original_filename else ""
            if extension not in {"pdf", "doc", "docx", "xls", "xlsx", "csv"}:
                flash("Bid documents must be PDF, Word, Excel, or CSV files.")
                return render_template("bid_submission_form.html", tender=tender, bidder=bidder), 400
            category = f"Bidder {bidder.id} · {document_category}"
            prior_versions = TenderDocument.query.filter_by(tender_id=tender.id, category=category, original_filename=original_filename).all()
            version = max((record.version for record in prior_versions), default=0) + 1
            stored_filename = f"{tender.id}_{uuid.uuid4().hex}_{original_filename}"
            upload_folder = tender_document_folder(tender.id)
            os.makedirs(upload_folder, exist_ok=True)
            upload.save(os.path.join(upload_folder, stored_filename))
            db.session.add(TenderDocument(
                tender=tender,
                category=category,
                original_filename=original_filename,
                stored_filename=stored_filename,
                version=version,
                description=request.form.get("document_title", "").strip() or None,
                uploaded_by=current_user,
            ))
            document_reference = stored_filename
        submission = BidSubmission(
            tender=tender,
            bidder=bidder,
            document_category=document_category,
            document_title=request.form.get("document_title"),
            document_reference=document_reference,
            status=request.form.get("status", "Received"),
        )
        db.session.add(submission)
        record_tender_event(tender, "BID RECEIPT", "Submission recorded", f"{document_category} recorded for {bidder.legal_name}.", current_user, document_reference)
        db.session.commit()
        log_audit("Tender Management", "create", "bid_submission", f"Recorded {submission.document_category} for {bidder.legal_name}.", user=current_user)
        flash("Bid submission record created successfully.")
        return redirect(url_for("tender.workspace", tender_id=tender.id))
    return render_template("bid_submission_form.html", tender=tender, bidder=bidder)


@tender_bp.route("/<int:tender_id>/compliance/new", methods=["GET", "POST"])
@login_required
def create_tender_compliance_item(tender_id):
    require_tender_admin()
    tender = Tender.query.get_or_404(tender_id)
    if request.method == "POST":
        item = TenderComplianceItem(
            tender=tender,
            requirement=request.form.get("requirement"),
            submitted=request.form.get("submitted") == "true",
            valid=None if request.form.get("valid") == "pending" else request.form.get("valid") == "true",
            notes=request.form.get("notes"),
            evidence_reference=request.form.get("evidence_reference"),
        )
        db.session.add(item)
        db.session.commit()
        log_audit("Tender Compliance", "create", "tender_compliance_item", f"Added compliance requirement for {tender.tender_number}.", user=current_user)
        flash("Compliance requirement created successfully.")
        return redirect(url_for("tender.workspace", tender_id=tender.id))
    return render_template("tender_compliance_form.html", tender=tender)


@tender_bp.route("/<int:tender_id>/criteria/new", methods=["GET", "POST"])
@login_required
def create_tender_criterion(tender_id):
    require_tender_admin()
    tender = Tender.query.get_or_404(tender_id)
    if request.method == "POST":
        if EvaluatorScore.query.filter_by(tender_id=tender.id, status="Submitted").first():
            flash("The evaluation methodology is locked after score sheets are submitted.")
            return redirect(url_for("tender.workspace", tender_id=tender.id)), 409
        name = request.form.get("name", "").strip()
        weight = float(request.form.get("weight", 0) or 0)
        if not name or weight <= 0 or sum(item.weight for item in tender.criteria) + weight > 100:
            flash("Enter a criterion name and positive weight; total criterion weight cannot exceed 100%.")
            return render_template("tender_criterion_form.html", tender=tender), 400
        criterion = TenderCriterion(
            tender=tender,
            name=name,
            description=request.form.get("description"),
            weight=weight,
            max_score=safe_int(request.form.get("max_score"), 5),
            scoring_framework=request.form.get("scoring_framework", "1-5"),
            pass_fail=request.form.get("pass_fail") == "true",
            instructions=request.form.get("instructions"),
        )
        db.session.add(criterion)
        criterion.config = TenderCriterionConfig(
            category=request.form.get("category", "Functionality").strip() or "Functionality",
            minimum_qualifying_score=float(request.form.get("minimum_qualifying_score") or 0) or None,
            mandatory=request.form.get("mandatory") == "true",
            evidence_required=request.form.get("evidence_required", "").strip() or None,
            position=len(tender.criteria),
        )
        record_tender_event(tender, "EVALUATION", "Methodology configured", f"Criterion {criterion.name} added.", current_user)
        db.session.commit()
        log_audit("Tender Management", "create", "tender_criterion", f"Added criterion {criterion.name} to {tender.tender_number}.", user=current_user)
        flash("Evaluation criterion created successfully.")
        return redirect(url_for("tender.workspace", tender_id=tender.id))
    return render_template("tender_criterion_form.html", tender=tender)


@tender_bp.route("/<int:tender_id>/criteria/<int:criterion_id>/edit", methods=["GET", "POST"])
@login_required
def edit_tender_criterion(tender_id, criterion_id):
    require_tender_admin()
    tender = Tender.query.get_or_404(tender_id)
    criterion = TenderCriterion.query.filter_by(id=criterion_id, tender_id=tender.id).first_or_404()
    if EvaluatorScore.query.filter_by(tender_id=tender.id, status="Submitted").first():
        flash("The evaluation methodology is locked after score sheets are submitted.")
        return redirect(url_for("tender.workspace", tender_id=tender.id)), 409
    if request.method == "POST":
        name = request.form.get("name", "").strip()
        weight = float(request.form.get("weight", 0) or 0)
        other_weight = sum(item.weight for item in tender.criteria if item.id != criterion.id)
        if not name or weight <= 0 or other_weight + weight > 100:
            flash("Enter a name and positive weight; total criterion weight cannot exceed 100%.")
            return render_template("tender_criterion_form.html", tender=tender, criterion=criterion), 400
        criterion.name = name
        criterion.description = request.form.get("description", "").strip() or None
        criterion.weight = weight
        criterion.max_score = float(request.form.get("max_score", 5) or 5)
        criterion.scoring_framework = request.form.get("scoring_framework", "1-5")
        criterion.pass_fail = request.form.get("pass_fail") == "true"
        criterion.instructions = request.form.get("instructions", "").strip() or None
        if not criterion.config:
            criterion.config = TenderCriterionConfig()
        criterion.config.category = request.form.get("category", "Functionality").strip() or "Functionality"
        criterion.config.minimum_qualifying_score = float(request.form.get("minimum_qualifying_score") or 0) or None
        criterion.config.mandatory = request.form.get("mandatory") == "true"
        criterion.config.evidence_required = request.form.get("evidence_required", "").strip() or None
        db.session.commit()
        record_tender_event(tender, "EVALUATION", "Methodology updated", f"Criterion {criterion.name} updated.", current_user)
        db.session.commit()
        log_audit("Tender Methodology", "update", "tender_criterion", f"Updated criterion {criterion.name} for {tender.tender_number}.", user=current_user)
        flash("Criterion updated.")
        return redirect(url_for("tender.workspace", tender_id=tender.id))
    return render_template("tender_criterion_form.html", tender=tender, criterion=criterion)


@tender_bp.route("/<int:tender_id>/criteria/<int:criterion_id>/delete", methods=["POST"])
@login_required
def delete_tender_criterion(tender_id, criterion_id):
    require_tender_admin()
    tender = Tender.query.get_or_404(tender_id)
    criterion = TenderCriterion.query.filter_by(id=criterion_id, tender_id=tender.id).first_or_404()
    if EvaluatorScore.query.filter_by(tender_id=tender.id, status="Submitted").first():
        flash("A criterion cannot be deleted after evaluation is locked.")
        return redirect(url_for("tender.workspace", tender_id=tender.id)), 409
    name = criterion.name
    db.session.delete(criterion)
    db.session.commit()
    record_tender_event(tender, "EVALUATION", "Methodology changed", f"Criterion {name} deleted before evaluation lock.", current_user)
    db.session.commit()
    log_audit("Tender Methodology", "delete", "tender_criterion", f"Deleted criterion {name} from {tender.tender_number}.", user=current_user)
    flash("Criterion deleted.")
    return redirect(url_for("tender.workspace", tender_id=tender.id))


@tender_bp.route("/<int:tender_id>/criteria/<int:criterion_id>/reorder", methods=["POST"])
@login_required
def reorder_tender_criterion(tender_id, criterion_id):
    require_tender_admin()
    tender = Tender.query.get_or_404(tender_id)
    if EvaluatorScore.query.filter_by(tender_id=tender.id, status="Submitted").first():
        return redirect(url_for("tender.workspace", tender_id=tender.id)), 409
    criterion = TenderCriterion.query.filter_by(id=criterion_id, tender_id=tender.id).first_or_404()
    ordered = sorted(tender.criteria, key=lambda item: item.config.position if item.config else item.id)
    index = ordered.index(criterion)
    target = index + (-1 if request.form.get("direction") == "up" else 1)
    if 0 <= target < len(ordered):
        ordered[index], ordered[target] = ordered[target], ordered[index]
    for position, item in enumerate(ordered):
        if not item.config:
            item.config = TenderCriterionConfig(position=position)
        else:
            item.config.position = position
    db.session.commit()
    log_audit("Tender Methodology", "reorder", "tender_criterion", f"Reordered criteria for {tender.tender_number}.", user=current_user)
    return redirect(url_for("tender.workspace", tender_id=tender.id))


@tender_bp.route("/<int:tender_id>/criteria/<int:criterion_id>/subcriteria/new", methods=["GET", "POST"])
@login_required
def create_tender_subcriterion(tender_id, criterion_id):
    require_tender_admin()
    tender = Tender.query.get_or_404(tender_id)
    criterion = TenderCriterion.query.filter_by(id=criterion_id, tender_id=tender.id).first_or_404()
    if request.method == "POST":
        if EvaluatorScore.query.filter_by(tender_id=tender.id, status="Submitted").first():
            flash("Subcriteria are locked after score sheets are submitted.")
            return redirect(url_for("tender.workspace", tender_id=tender.id)), 409
        name = request.form.get("name", "").strip()
        max_points = float(request.form.get("max_points") or 0)
        current_total = sum(item.max_points for item in criterion.subcriteria)
        if not name or max_points <= 0 or current_total + max_points > criterion.max_score:
            flash("Enter a name and positive points; subcriterion points cannot exceed the parent maximum.")
            return render_template("tender_subcriterion_form.html", tender=tender, criterion=criterion), 400
        subcriterion = TenderSubcriterion(
            criterion=criterion,
            name=name,
            description=request.form.get("description", "").strip() or None,
            max_points=max_points,
            position=len(criterion.subcriteria),
        )
        db.session.add(subcriterion)
        db.session.commit()
        log_audit("Tender Methodology", "create", "tender_subcriterion", f"Added {name} under {criterion.name} for {tender.tender_number}.", user=current_user)
        flash("Subcriterion added successfully.")
        return redirect(url_for("tender.workspace", tender_id=tender.id))
    return render_template("tender_subcriterion_form.html", tender=tender, criterion=criterion)


@tender_bp.route("/<int:tender_id>/criteria/<int:criterion_id>/subcriteria/<int:subcriterion_id>/edit", methods=["GET", "POST"])
@login_required
def edit_tender_subcriterion(tender_id, criterion_id, subcriterion_id):
    require_tender_admin()
    tender = Tender.query.get_or_404(tender_id)
    criterion = TenderCriterion.query.filter_by(id=criterion_id, tender_id=tender.id).first_or_404()
    subcriterion = TenderSubcriterion.query.filter_by(id=subcriterion_id, criterion_id=criterion.id).first_or_404()
    if EvaluatorScore.query.filter_by(tender_id=tender.id, status="Submitted").first():
        flash("Subcriteria are locked after score sheets are submitted.")
        return redirect(url_for("tender.workspace", tender_id=tender.id)), 409
    if request.method == "POST":
        name = request.form.get("name", "").strip()
        max_points = float(request.form.get("max_points") or 0)
        other_points = sum(item.max_points for item in criterion.subcriteria if item.id != subcriterion.id)
        if not name or max_points <= 0 or other_points + max_points > criterion.max_score:
            flash("Enter a valid name and points; subcriterion points cannot exceed the parent maximum.")
            return render_template("tender_subcriterion_form.html", tender=tender, criterion=criterion, subcriterion=subcriterion), 400
        subcriterion.name = name
        subcriterion.max_points = max_points
        subcriterion.description = request.form.get("description", "").strip() or None
        db.session.commit()
        log_audit("Tender Methodology", "update", "tender_subcriterion", f"Updated subcriterion {name} for {tender.tender_number}.", user=current_user)
        return redirect(url_for("tender.workspace", tender_id=tender.id))
    return render_template("tender_subcriterion_form.html", tender=tender, criterion=criterion, subcriterion=subcriterion)


@tender_bp.route("/<int:tender_id>/criteria/<int:criterion_id>/subcriteria/<int:subcriterion_id>/delete", methods=["POST"])
@login_required
def delete_tender_subcriterion(tender_id, criterion_id, subcriterion_id):
    require_tender_admin()
    tender = Tender.query.get_or_404(tender_id)
    criterion = TenderCriterion.query.filter_by(id=criterion_id, tender_id=tender.id).first_or_404()
    subcriterion = TenderSubcriterion.query.filter_by(id=subcriterion_id, criterion_id=criterion.id).first_or_404()
    if EvaluatorScore.query.filter_by(tender_id=tender.id, status="Submitted").first():
        return redirect(url_for("tender.workspace", tender_id=tender.id)), 409
    name = subcriterion.name
    db.session.delete(subcriterion)
    db.session.commit()
    log_audit("Tender Methodology", "delete", "tender_subcriterion", f"Deleted subcriterion {name} from {tender.tender_number}.", user=current_user)
    return redirect(url_for("tender.workspace", tender_id=tender.id))


@tender_bp.route("/<int:tender_id>/evaluators", methods=["POST"])
@login_required
def assign_tender_evaluator(tender_id):
    require_tender_admin()
    tender = Tender.query.get_or_404(tender_id)
    user = User.query.filter_by(id=safe_int(request.form.get("user_id")), is_active=True).first()
    role_name = user.role_obj.name if user and user.role_obj else None
    if not user or role_name not in TENDER_EVALUATOR_ROLES:
        flash("Select an active SCM or BEC evaluator account.")
        return redirect(url_for("tender.workspace", tender_id=tender.id)), 400
    assignment = TenderEvaluatorAssignment.query.filter_by(tender_id=tender.id, user_id=user.id).first()
    if assignment:
        assignment.active = True
        assignment.assigned_by = current_user
        assignment.assigned_at = datetime.utcnow()
    else:
        db.session.add(TenderEvaluatorAssignment(tender=tender, user=user, assigned_by=current_user))
    record_tender_event(tender, "EVALUATION", "Evaluator assigned", f"{user.full_name} assigned to the evaluation.", current_user)
    db.session.commit()
    log_audit("Tender Evaluation", "assign_evaluator", "tender_evaluator", f"Assigned {user.full_name} to {tender.tender_number}.", user=current_user)
    flash(f"{user.full_name} assigned as a tender evaluator.")
    return redirect(url_for("tender.workspace", tender_id=tender.id))


@tender_bp.route("/<int:tender_id>/screening", methods=["GET", "POST"])
@login_required
def tender_screening(tender_id):
    require_tender_admin()
    tender = Tender.query.get_or_404(tender_id)
    if request.method == "POST":
        action = request.form.get("action")
        if action == "add_requirement":
            name = request.form.get("name", "").strip()
            if not name:
                flash("A screening requirement is required.")
                return redirect(url_for("tender.tender_screening", tender_id=tender.id)), 400
            db.session.add(TenderScreeningRequirement(
                tender=tender,
                name=name,
                category=request.form.get("category", "Mandatory").strip() or "Mandatory",
                mandatory=request.form.get("mandatory") == "true",
                position=len(tender.screening_requirements),
            ))
            record_tender_event(tender, "SCREENING", "Requirement added", f"Screening requirement added: {name}.", current_user)
            db.session.commit()
            log_audit("Tender Screening", "add_requirement", "screening_requirement", f"Added screening requirement for {tender.tender_number}: {name}.", user=current_user)
            flash("Screening requirement added.")
        else:
            requirement = TenderScreeningRequirement.query.filter_by(
                id=safe_int(request.form.get("requirement_id")), tender_id=tender.id
            ).first_or_404()
            bidder = TenderBidder.query.filter_by(
                id=safe_int(request.form.get("bidder_id")), tender_id=tender.id
            ).first_or_404()
            result_value = request.form.get("result", "PENDING").upper()
            if result_value not in {"PASS", "FAIL", "PENDING", "CLARIFICATION"}:
                abort(400)
            result = TenderScreeningResult.query.filter_by(requirement_id=requirement.id, bidder_id=bidder.id).first()
            if not result:
                result = TenderScreeningResult(requirement=requirement, bidder=bidder)
                db.session.add(result)
            result.result = result_value
            result.evidence_reference = request.form.get("evidence_reference", "").strip() or None
            result.comments = request.form.get("comments", "").strip() or None
            result.reviewer = current_user
            result.reviewed_at = datetime.utcnow()
            record_tender_event(tender, "SCREENING", result_value, f"{bidder.legal_name}: {requirement.name}.", current_user, result.evidence_reference)
            db.session.commit()
            log_audit("Tender Screening", "review", "screening_result", f"{bidder.legal_name}: {requirement.name} marked {result_value}.", user=current_user)
            flash("Bid screening result saved.")
        return redirect(url_for("tender.tender_screening", tender_id=tender.id))
    screening_results = {
        (result.requirement_id, result.bidder_id): result
        for requirement in tender.screening_requirements for result in requirement.results
    }
    return render_template("tender_screening.html", tender=tender, screening_results=screening_results)


@tender_bp.route("/<int:tender_id>/evaluations/new", methods=["GET", "POST"])
@login_required
def create_evaluator_score(tender_id):
    tender = Tender.query.get_or_404(tender_id)
    require_tender_evaluator(tender)
    bidders = tender.bidders
    criteria = tender.criteria
    if request.method == "POST":
        bidder_id = safe_int(request.form.get("bidder_id"))
        criterion_id = safe_int(request.form.get("criterion_id"))
        bidder = TenderBidder.query.filter_by(id=bidder_id, tender_id=tender.id).first()
        criterion = TenderCriterion.query.filter_by(id=criterion_id, tender_id=tender.id).first()
        if not bidder or not criterion:
            flash("Select a valid bidder and criterion for this tender.")
            return render_template("evaluator_score_form.html", tender=tender, bidders=bidders, criteria=criteria), 400
        for requirement in tender.screening_requirements:
            screening_result = TenderScreeningResult.query.filter_by(
                requirement_id=requirement.id, bidder_id=bidder.id
            ).first()
            if requirement.mandatory and (not screening_result or screening_result.result != "PASS"):
                flash(f"Mandatory screening is not passed for {bidder.legal_name}: {requirement.name}.")
                return redirect(url_for("tender.tender_screening", tender_id=tender.id)), 409
        evaluation = EvaluatorScore.query.filter_by(tender_id=tender.id, bidder_id=bidder.id, criterion_id=criterion.id, evaluator_id=current_user.id).first()
        if evaluation and evaluation.status == "Submitted":
            flash("This evaluation is locked after submission. An authorised reviewer must reopen it with a reason.")
            return redirect(url_for("tender.workspace", tender_id=tender.id)), 409
        criterion_subcriteria = criterion.subcriteria
        child_values = {}
        if criterion_subcriteria:
            for subcriterion in criterion_subcriteria:
                raw_score = request.form.get(f"subcriterion_{subcriterion.id}", "").strip()
                try:
                    child_score = float(raw_score)
                except ValueError:
                    flash(f"Enter a score for {subcriterion.name}.")
                    return render_template("evaluator_score_form.html", tender=tender, bidders=bidders, criteria=criteria), 400
                if not math.isfinite(child_score) or not 0 <= child_score <= subcriterion.max_points:
                    flash(f"The score for {subcriterion.name} must be between 0 and {subcriterion.max_points:g}.")
                    return render_template("evaluator_score_form.html", tender=tender, bidders=bidders, criteria=criteria), 400
                child_values[subcriterion.id] = child_score
            score = sum(child_values.values())
            if score > criterion.max_score:
                flash("The subcriterion total cannot exceed the parent criterion maximum.")
                return render_template("evaluator_score_form.html", tender=tender, bidders=bidders, criteria=criteria), 400
        elif criterion.pass_fail or criterion.scoring_framework.strip().lower() == "pass/fail":
            pass_fail_value = request.form.get("pass_fail_score", "")
            if pass_fail_value not in {"Pass", "Fail"}:
                flash("Select Pass or Fail for this criterion.")
                return render_template("evaluator_score_form.html", tender=tender, bidders=bidders, criteria=criteria), 400
            score = criterion.max_score if pass_fail_value == "Pass" else 0
        else:
            try:
                score = float(request.form.get("score", ""))
            except ValueError:
                flash("Enter a valid numeric score.")
                return render_template("evaluator_score_form.html", tender=tender, bidders=bidders, criteria=criteria), 400
            if not math.isfinite(score) or not 0 <= score <= criterion.max_score:
                flash(f"Score must be between 0 and {criterion.max_score}.")
                return render_template("evaluator_score_form.html", tender=tender, bidders=bidders, criteria=criteria), 400
        if not evaluation:
            evaluation = EvaluatorScore(tender=tender, bidder=bidder, criterion=criterion, evaluator=current_user)
            db.session.add(evaluation)
            db.session.flush()
        evaluation.score = score
        evaluation.comments = request.form.get("comments")
        evaluation.status = "Submitted" if request.form.get("action") == "submit" else "Draft"
        evaluation.submitted_at = datetime.utcnow() if evaluation.status == "Submitted" else None
        evaluation.locked_at = datetime.utcnow() if evaluation.status == "Submitted" else None
        for subcriterion_id, child_score in child_values.items():
            child_result = EvaluatorSubcriterionScore.query.filter_by(
                evaluation_id=evaluation.id, subcriterion_id=subcriterion_id
            ).first()
            if not child_result:
                child_result = EvaluatorSubcriterionScore(evaluation=evaluation, subcriterion_id=subcriterion_id, score=child_score)
                db.session.add(child_result)
            child_result.score = child_score
            child_result.comments = request.form.get(f"subcriterion_comment_{subcriterion_id}", "").strip() or None
            child_result.evidence_reference = request.form.get(f"subcriterion_evidence_{subcriterion_id}", "").strip() or None
        record_tender_event(tender, "EVALUATION", evaluation.status, f"{current_user.full_name} {evaluation.status.lower()} {bidder.legal_name} / {criterion.name}.", current_user)
        db.session.commit()
        log_audit("Tender Evaluation", evaluation.status.lower(), "evaluator_score", f"{evaluation.status} score for {bidder.legal_name} / {criterion.name}.", user=current_user)
        flash(f"Evaluation {evaluation.status.lower()} successfully.")
        return redirect(url_for("tender.workspace", tender_id=tender.id))
    return render_template("evaluator_score_form.html", tender=tender, bidders=bidders, criteria=criteria)


@tender_bp.route("/<int:tender_id>/evaluations/<int:evaluation_id>/reopen", methods=["POST"])
@login_required
def reopen_evaluator_score(tender_id, evaluation_id):
    if not current_user.role_obj or current_user.role_obj.name not in {"System Administrator", "SCM Manager", "Internal Audit"}:
        abort(403)
    evaluation = EvaluatorScore.query.filter_by(id=evaluation_id, tender_id=tender_id).first_or_404()
    reason = request.form.get("reason", "").strip()
    if not reason:
        flash("A reason is required to reopen a submitted evaluation.")
        return redirect(url_for("tender.workspace", tender_id=tender_id)), 400
    prior_subcriteria = [
        {"subcriterion_id": result.subcriterion_id, "score": result.score, "comments": result.comments, "evidence_reference": result.evidence_reference}
        for result in evaluation.subcriterion_scores
    ]
    db.session.add(EvaluatorScoreRevision(
        evaluation=evaluation,
        reopened_by=current_user,
        reason=reason,
        previous_score=evaluation.score,
        previous_comments=evaluation.comments,
        previous_status=evaluation.status,
        previous_subcriteria=json.dumps(prior_subcriteria),
    ))
    evaluation.status = "Draft"
    evaluation.reopened_at = datetime.utcnow()
    evaluation.reopen_reason = reason
    evaluation.locked_at = None
    db.session.commit()
    record_tender_event(evaluation.tender, "EVALUATION", "Reopened", f"Score sheet {evaluation.id} reopened: {reason}", current_user)
    db.session.commit()
    log_audit("Tender Evaluation", "reopen", "evaluator_score", f"Reopened evaluation {evaluation.id}: {reason}", user=current_user)
    flash("Evaluation reopened under controlled review.")
    return redirect(url_for("tender.workspace", tender_id=tender_id))


bec_bp = Blueprint("bec", __name__, url_prefix="/bec")


@bec_bp.route("/")
@login_required
def index():
    role_name = current_user.role_obj.name if current_user.role_obj else None
    if role_name in TENDER_ADMIN_ROLES:
        items = BidEvaluationCommittee.query.order_by(BidEvaluationCommittee.created_at.desc()).all()
    elif role_name == "BEC Member":
        assigned_ids = db.session.query(TenderEvaluatorAssignment.tender_id).filter_by(user_id=current_user.id, active=True)
        items = BidEvaluationCommittee.query.filter(BidEvaluationCommittee.tender_id.in_(assigned_ids)).order_by(BidEvaluationCommittee.created_at.desc()).all()
    else:
        abort(403)
    return render_template("module_list.html", title="BEC", items=items, columns=["Tender", "Status", "Evaluation Status"])


@bec_bp.route("/new", methods=["GET", "POST"])
@login_required
def create_bec_record():
    require_tender_admin()
    tenders = Tender.query.order_by(Tender.tender_number).all()
    if request.method == "POST":
        tender_id = safe_int(request.form.get("tender_id"))
        tender = Tender.query.get(tender_id) if tender_id is not None else None
        if not tender:
            flash("Select a valid tender for the BEC record.")
            return render_template("bec_form.html", tenders=tenders), 400
        record = BidEvaluationCommittee(
            tender=tender,
            title=request.form.get("title"),
            evaluation_status=request.form.get("evaluation_status", "In Progress"),
        )
        db.session.add(record)
        db.session.commit()
        log_audit("BEC", "create", "bid_evaluation_committee", f"Created BEC record {record.title}.", user=current_user)
        flash("BEC record created successfully.")
        return redirect(url_for("bec.index"))
    return render_template("bec_form.html", tenders=tenders)


bac_bp = Blueprint("bac", __name__, url_prefix="/bac")


@bac_bp.route("/")
@login_required
def index():
    require_tender_admin()
    items = BidAdjudicationCommittee.query.order_by(BidAdjudicationCommittee.created_at.desc()).all()
    return render_template("module_list.html", title="BAC", items=items, columns=["Tender", "Decision", "Status"])


@bac_bp.route("/new", methods=["GET", "POST"])
@login_required
def create_bac_record():
    require_tender_admin()
    tenders = Tender.query.order_by(Tender.tender_number).all()
    if request.method == "POST":
        tender_id = safe_int(request.form.get("tender_id"))
        tender = Tender.query.get(tender_id) if tender_id is not None else None
        if not tender:
            flash("Select a valid tender for the BAC record.")
            return render_template("bac_form.html", tenders=tenders), 400
        record = BidAdjudicationCommittee(
            tender=tender,
            title=request.form.get("title"),
            decision=request.form.get("decision"),
            status=request.form.get("status", "Pending"),
        )
        db.session.add(record)
        db.session.commit()
        log_audit("BAC", "create", "bid_adjudication_committee", f"Created BAC record {record.title}.", user=current_user)
        flash("BAC record created successfully.")
        return redirect(url_for("bac.index"))
    return render_template("bac_form.html", tenders=tenders)


# Contracts
contract_bp = Blueprint("contract", __name__, url_prefix="/contract")


@contract_bp.route("/")
@login_required
def index():
    status = request.args.get("status")
    query = Contract.query
    if status:
        query = query.filter_by(status=status)
    items = query.order_by(Contract.created_at.desc()).all()
    return render_template("module_list.html", title="Contracts", items=items, columns=["Contract Number", "Supplier", "Value", "Status"], active_filter=status, empty_message=f"No {status.lower()} contracts" if status else "No contracts here yet")


@contract_bp.route("/new", methods=["GET", "POST"])
@login_required
def create_contract():
    if request.method == "POST":
        contract = Contract(
            contract_number=request.form.get("contract_number"),
            supplier=request.form.get("supplier"),
            tender=request.form.get("tender"),
            start_date=request.form.get("start_date"),
            end_date=request.form.get("end_date"),
            value=float(request.form.get("value", 0) or 0),
            status=request.form.get("status", "Active"),
        )
        db.session.add(contract)
        db.session.commit()
        log_audit("Contracts", "create", "contract", f"Created contract {contract.contract_number}.", user=current_user)
        flash("Contract created successfully.")
        return redirect(url_for("contract.index"))
    return render_template("contract_form.html")


# Transport and travel
transport_bp = Blueprint("transport", __name__, url_prefix="/transport")

DIRECTORATE_TRANSPORT_ROLES = {"System Administrator", "Directorate DD", "Directorate Director", "Directorate Chief Director"}
CENTRAL_TRANSPORT_ROLES = {"System Administrator", "Transport Manager", "Transport Supervisor", "Transport Officer", "Fleet Administrator"}


def require_central_transport_role():
    if not current_user.role_obj or current_user.role_obj.name not in CENTRAL_TRANSPORT_ROLES:
        abort(403)


def require_directorate_transport_role():
    if not current_user.role_obj or current_user.role_obj.name not in DIRECTORATE_TRANSPORT_ROLES:
        abort(403)
    if current_user.role_obj.name != "System Administrator" and not current_user.employee:
        abort(403)


def record_transport_event(transport_request, action, new_status=None, comment=None):
    db.session.add(TransportCaseEvent(
        transport_request=transport_request,
        action=action,
        previous_status=transport_request.status,
        new_status=new_status or transport_request.status,
        comment=comment,
        user=current_user,
        role_name=current_user.role_obj.name if current_user.role_obj else None,
    ))


@transport_bp.route("/")
@login_required
def index():
    require_central_transport_role()
    items = TransportPortalRequest.query.order_by(TransportPortalRequest.created_at.desc()).all()
    allocations = TransportAllocation.query.order_by(TransportAllocation.created_at.desc()).all()
    return render_template(
        "transport_command.html",
        items=items,
        allocations=allocations,
        stats={
            "intake": sum(item.status in {"Submitted", "Received by Central SCM"} for item in items),
            "active": sum(item.status not in {"Completed", "Closed"} for item in items),
            "vehicles": Vehicle.query.count(),
            "available_vehicles": Vehicle.query.filter_by(availability_status="AVAILABLE").count(),
            "drivers": Driver.query.count(),
            "active_trips": sum(item.status in {"Vehicle Allocated", "AUTHORISED", "ISSUED", "STARTED", "IN_PROGRESS"} for item in items),
            "maintenance": TransportMaintenance.query.filter_by(status="OPEN").count(),
            "incidents": TransportIncident.query.filter_by(status="OPEN").count(),
        },
    )


@transport_bp.route("/requests/<int:request_id>")
@login_required
def transport_case(request_id):
    require_central_transport_role()
    item = TransportPortalRequest.query.get_or_404(request_id)
    return render_template("transport_case.html", item=item, events=item.case_events, documents=item.transport_documents)


@transport_bp.route("/requests/<int:request_id>/status", methods=["POST"])
@login_required
def update_transport_case_status(request_id):
    require_central_transport_role()
    item = TransportPortalRequest.query.get_or_404(request_id)
    new_status = request.form.get("status", "").strip()
    allowed_statuses = {"Received by Central SCM", "Validation", "Correction Required", "Assigned", "Vehicle Allocated", "AUTHORISED", "ISSUED", "IN_PROGRESS", "Completed", "Closed"}
    if new_status not in allowed_statuses:
        abort(400)
    comment = request.form.get("comment", "").strip() or None
    item.status = new_status
    record_transport_event(item, "Status updated", new_status, comment)
    db.session.commit()
    log_audit("SmartFleet", "status_update", "transport_request", f"{item.id} moved to {new_status}.", user=current_user)
    flash(f"Transport case {item.id} moved to {new_status}.")
    return redirect(url_for("transport.transport_case", request_id=item.id))


@transport_bp.route("/requests/<int:request_id>/assign", methods=["POST"])
@login_required
def assign_transport_case(request_id):
    require_central_transport_role()
    item = TransportPortalRequest.query.get_or_404(request_id)
    official = User.query.filter_by(id=safe_int(request.form.get("official_id")), is_active=True).first()
    if not official or not official.role_obj or official.role_obj.name not in CENTRAL_TRANSPORT_ROLES:
        abort(400)
    assignment = item.transport_assignment or TransportAssignment(transport_request=item)
    assignment.official = official
    assignment.unit = "Central Transport SCM"
    assignment.stage = request.form.get("stage", "Central Intake")
    assignment.assigned_by = current_user
    db.session.add(assignment)
    record_transport_event(item, "Case assigned", item.status, f"Assigned to {official.full_name} for {assignment.stage}.")
    db.session.commit()
    log_audit("SmartFleet", "assign", "transport_request", f"Assigned case {item.id} to {official.username}.", user=current_user)
    flash("Transport case assigned successfully.")
    return redirect(url_for("transport.transport_case", request_id=item.id))


@transport_bp.route("/new", methods=["GET", "POST"])
@login_required
def create_trip_request():
    if request.method == "POST":
        trip = TripRequest(
            title=request.form.get("title"),
            origin=request.form.get("origin"),
            destination=request.form.get("destination"),
            purpose=request.form.get("purpose"),
            status=request.form.get("status", "Requested"),
        )
        db.session.add(trip)
        db.session.commit()
        log_audit("Transport", "create", "trip_request", f"Created trip request {trip.title}.", user=current_user)
        flash("Trip request created successfully.")
        return redirect(url_for("transport.index"))
    return render_template("trip_request_form.html")


TRANSPORT_ROLES = {"System Administrator", "Transport Manager", "Transport Officer"}


def require_transport_role():
    require_central_transport_role()


@transport_bp.route("/vehicles")
@login_required
def vehicles():
    require_transport_role()
    items = Vehicle.query.order_by(Vehicle.fleet_number).all()
    return render_template("vehicle_list.html", items=items)


@transport_bp.route("/vehicles/import-template")
@login_required
def vehicle_import_template():
    require_transport_role()
    columns = "registration_number,fleet_number,make,model,vehicle_type,fuel_type,fuel_capacity,current_odometer,current_location,next_service,condition,notes\n"
    return Response(
        columns,
        mimetype="text/csv",
        headers={"Content-Disposition": "attachment; filename=vehicle-import-template.csv"},
    )


@transport_bp.route("/vehicles/import", methods=["POST"])
@login_required
def import_vehicles():
    require_transport_role()
    upload = request.files.get("vehicle_csv")
    if not upload or not upload.filename:
        flash("Choose a CSV file to import.")
        return redirect(url_for("transport.vehicles"))
    if not upload.filename.lower().endswith(".csv"):
        flash("Vehicle imports must be CSV files.")
        return redirect(url_for("transport.vehicles"))

    try:
        content = upload.stream.read().decode("utf-8-sig")
    except UnicodeDecodeError:
        flash("The CSV file must use UTF-8 encoding.")
        return redirect(url_for("transport.vehicles"))

    reader = csv.DictReader(StringIO(content))
    if not reader.fieldnames:
        flash("The CSV file is empty or missing its header row.")
        return redirect(url_for("transport.vehicles"))

    required_columns = {"registration_number", "fleet_number"}
    columns = {column.strip().lower() for column in reader.fieldnames if column}
    missing_columns = required_columns - columns
    if missing_columns:
        flash("Missing required CSV columns: " + ", ".join(sorted(missing_columns)) + ".")
        return redirect(url_for("transport.vehicles"))

    registrations = {value for (value,) in db.session.query(Vehicle.registration_number).all()}
    fleet_numbers = {value for (value,) in db.session.query(Vehicle.fleet_number).all()}
    vehicles_to_add = []
    errors = []
    for row_number, raw_row in enumerate(reader, start=2):
        row = {key.strip().lower(): (value or "").strip() for key, value in raw_row.items() if key}
        registration_number = row.get("registration_number", "")
        fleet_number = row.get("fleet_number", "")
        row_errors = []
        if not registration_number or not fleet_number:
            row_errors.append("registration_number and fleet_number are required")
        if registration_number and registration_number in registrations:
            row_errors.append(f"registration number {registration_number} already exists")
        if fleet_number and fleet_number in fleet_numbers:
            row_errors.append(f"fleet number {fleet_number} already exists")
        if registration_number:
            registrations.add(registration_number)
        if fleet_number:
            fleet_numbers.add(fleet_number)

        try:
            fuel_capacity = float(row.get("fuel_capacity") or 0)
            if not math.isfinite(fuel_capacity) or fuel_capacity < 0:
                raise ValueError
        except ValueError:
            fuel_capacity = 0
            row_errors.append("fuel_capacity must be a non-negative number")
        try:
            current_odometer = int(row.get("current_odometer") or 0)
            if current_odometer < 0:
                raise ValueError
        except ValueError:
            current_odometer = 0
            row_errors.append("current_odometer must be a non-negative whole number")

        if row_errors:
            errors.append(f"Row {row_number}: " + "; ".join(row_errors))
            continue

        vehicles_to_add.append(
            Vehicle(
                registration_number=registration_number,
                fleet_number=fleet_number,
                make=row.get("make") or None,
                model=row.get("model") or None,
                vehicle_type=row.get("vehicle_type") or None,
                fuel_type=row.get("fuel_type") or None,
                fuel_capacity=fuel_capacity,
                current_odometer=current_odometer,
                current_location=row.get("current_location") or None,
                next_service=row.get("next_service") or None,
                condition=row.get("condition") or "Good",
                notes=row.get("notes") or None,
            )
        )

    if not vehicles_to_add and not errors:
        errors.append("The CSV contains no vehicle rows.")
    if errors:
        flash("Import canceled. " + " | ".join(errors[:5]))
        return redirect(url_for("transport.vehicles"))

    try:
        db.session.add_all(vehicles_to_add)
        db.session.commit()
    except IntegrityError:
        db.session.rollback()
        flash("Import canceled because a registration or fleet number already exists.")
        return redirect(url_for("transport.vehicles"))

    log_audit("Transport", "import", "vehicle", f"Imported {len(vehicles_to_add)} fleet vehicles.", user=current_user)
    flash(f"Imported {len(vehicles_to_add)} vehicles successfully.")
    return redirect(url_for("transport.vehicles"))


@transport_bp.route("/vehicles/new", methods=["GET", "POST"])
@login_required
def create_vehicle():
    require_transport_role()
    if request.method == "POST":
        vehicle = Vehicle(registration_number=request.form.get("registration_number"), fleet_number=request.form.get("fleet_number"), make=request.form.get("make"), model=request.form.get("model"), vehicle_type=request.form.get("vehicle_type"), fuel_type=request.form.get("fuel_type"), fuel_capacity=float(request.form.get("fuel_capacity", 0) or 0), current_odometer=safe_int(request.form.get("current_odometer"), 0), current_location=request.form.get("current_location"), availability_status="AVAILABLE", next_service=request.form.get("next_service"), condition=request.form.get("condition", "Good"), notes=request.form.get("notes"))
        db.session.add(vehicle)
        db.session.commit()
        log_audit("Transport", "create", "vehicle", f"Created vehicle {vehicle.fleet_number}.", user=current_user)
        flash("Vehicle created successfully.")
        return redirect(url_for("transport.vehicles"))
    return render_template("vehicle_form.html")


@transport_bp.route("/drivers/new", methods=["GET", "POST"])
@login_required
def create_driver():
    require_transport_role()
    if request.method == "POST":
        driver = Driver(employee_name=request.form.get("employee_name"), employee_number=request.form.get("employee_number"), contact=request.form.get("contact"), licence_number=request.form.get("licence_number"), licence_expiry=request.form.get("licence_expiry"), status=request.form.get("status", "AUTHORISED"))
        db.session.add(driver)
        db.session.commit()
        log_audit("Transport", "create", "driver", f"Created driver record for {driver.employee_name}.", user=current_user)
        flash("Driver created successfully.")
        return redirect(url_for("transport.vehicles"))
    return render_template("driver_form.html")


@transport_bp.route("/requests/<int:request_id>/allocate", methods=["GET", "POST"])
@login_required
def allocate_transport_request(request_id):
    require_central_transport_role()
    transport_request = TransportPortalRequest.query.get_or_404(request_id)
    vehicles_available = Vehicle.query.filter_by(availability_status="AVAILABLE").order_by(Vehicle.fleet_number).all()
    drivers = Driver.query.filter_by(status="AUTHORISED").order_by(Driver.employee_name).all()
    if request.method == "POST":
        vehicle = Vehicle.query.get(safe_int(request.form.get("vehicle_id")))
        driver = Driver.query.get(safe_int(request.form.get("driver_id"))) if request.form.get("driver_id") else None
        if not vehicle or vehicle.availability_status != "AVAILABLE":
            flash("Select an available vehicle.")
            return render_template("transport_allocation_form.html", item=transport_request, vehicles=vehicles_available, drivers=drivers), 400
        if TransportMaintenance.query.filter_by(vehicle_id=vehicle.id, status="OPEN").first():
            flash("This vehicle is under maintenance and cannot be allocated.")
            return render_template("transport_allocation_form.html", item=transport_request, vehicles=vehicles_available, drivers=drivers), 409
        if driver and driver.status != "AUTHORISED":
            flash("Select an authorised driver.")
            return render_template("transport_allocation_form.html", item=transport_request, vehicles=vehicles_available, drivers=drivers), 400
        allocation = TransportAllocation(transport_request=transport_request, vehicle=vehicle, driver=driver, status="ALLOCATED", odometer_start=vehicle.current_odometer)
        vehicle.availability_status = "ALLOCATED"
        transport_request.status = "Vehicle Allocated"
        db.session.add(allocation)
        record_transport_event(transport_request, "Vehicle and driver allocated", "Vehicle Allocated", f"Vehicle {vehicle.fleet_number}; driver {driver.employee_name if driver else 'not assigned'}.")
        db.session.commit()
        log_audit("Transport", "allocate", "transport_allocation", f"Allocated {vehicle.fleet_number} to request {transport_request.id}.", user=current_user)
        flash("Vehicle allocated successfully.")
        return redirect(url_for("transport.index"))
    return render_template("transport_allocation_form.html", item=transport_request, vehicles=vehicles_available, drivers=drivers)


@transport_bp.route("/requests/<int:request_id>/authorise", methods=["POST"])
@login_required
def authorise_transport_trip(request_id):
    require_central_transport_role()
    transport_request = TransportPortalRequest.query.get_or_404(request_id)
    allocation = transport_request.allocation
    if not allocation or not allocation.driver:
        flash("Vehicle and driver allocation are required before Trip Authorisation.")
        return redirect(url_for("transport.transport_case", request_id=request_id)), 409
    version = (db.session.query(db.func.max(TransportTripAuthorisation.version)).filter_by(transport_request_id=request_id).scalar() or 0) + 1
    authorisation = TransportTripAuthorisation(transport_request=transport_request, allocation=allocation, authorised_by=current_user, version=version)
    db.session.add(authorisation)
    transport_request.status = "AUTHORISED"
    record_transport_event(transport_request, "Trip Authorisation generated", "AUTHORISED", f"Trip Authorisation version {version} issued.")
    db.session.commit()
    log_audit("SmartFleet", "trip_authorisation", "transport_request", f"Generated Trip Authorisation v{version} for case {request_id}.", user=current_user)
    flash("Trip Authorisation generated.")
    return redirect(url_for("transport.transport_case", request_id=request_id))


@transport_bp.route("/allocations/<int:allocation_id>/issue", methods=["POST"])
@login_required
def issue_transport_vehicle(allocation_id):
    require_central_transport_role()
    allocation = TransportAllocation.query.get_or_404(allocation_id)
    odometer = safe_int(request.form.get("odometer"))
    if not allocation.transport_request.trip_authorisations:
        flash("Generate Trip Authorisation before issuing the vehicle.")
        return redirect(url_for("transport.transport_case", request_id=allocation.transport_request.id)), 409
    if odometer is None or allocation.odometer_start is None or odometer < allocation.odometer_start:
        flash("Issue odometer must be greater than or equal to the allocation odometer.")
        return redirect(url_for("transport.transport_case", request_id=allocation.transport_request.id)), 400
    issue = allocation.vehicle_issue or TransportVehicleIssue(allocation=allocation, issued_by=current_user)
    issue.odometer = odometer
    issue.fuel_level = float(request.form.get("fuel_level") or 0)
    issue.condition = request.form.get("condition")
    issue.accessories = request.form.get("accessories")
    issue.acknowledgement = request.form.get("acknowledgement")
    db.session.add(issue)
    allocation.status = "ISSUED"
    allocation.transport_request.status = "ISSUED"
    record_transport_event(allocation.transport_request, "Vehicle issued", "ISSUED", f"Vehicle issued at odometer {odometer}.")
    db.session.commit()
    flash("Vehicle issuing record saved.")
    return redirect(url_for("transport.transport_case", request_id=allocation.transport_request.id))


@transport_bp.route("/allocations/<int:allocation_id>/start", methods=["POST"])
@login_required
def start_transport_trip(allocation_id):
    require_central_transport_role()
    allocation = TransportAllocation.query.get_or_404(allocation_id)
    if not allocation.vehicle_issue:
        flash("Issue the vehicle before starting the trip.")
        return redirect(url_for("transport.transport_case", request_id=allocation.transport_request.id)), 409
    allocation.status = "IN_PROGRESS"
    allocation.transport_request.status = "IN_PROGRESS"
    record_transport_event(allocation.transport_request, "Trip started", "IN_PROGRESS", "Vehicle departed under the authorised trip record.")
    db.session.commit()
    flash("Trip started.")
    return redirect(url_for("transport.transport_case", request_id=allocation.transport_request.id))


@transport_bp.route("/allocations/<int:allocation_id>/complete", methods=["POST"])
@login_required
def complete_transport_allocation(allocation_id):
    require_central_transport_role()
    allocation = TransportAllocation.query.get_or_404(allocation_id)
    odometer_end = safe_int(request.form.get("odometer_end"))
    if odometer_end is None or allocation.odometer_start is None or odometer_end < allocation.odometer_start:
        flash("A valid closing odometer greater than or equal to the opening reading is required.")
        return redirect(url_for("transport.index")), 400
    allocation.odometer_end = odometer_end
    allocation.fuel_issued = float(request.form.get("fuel_issued", 0) or 0)
    allocation.fuel_cost = float(request.form.get("fuel_cost", 0) or 0)
    allocation.condition_notes = request.form.get("condition_notes")
    allocation.status = "COMPLETED"
    allocation.returned_at = datetime.utcnow()
    allocation.vehicle.current_odometer = odometer_end
    allocation.vehicle.availability_status = "AVAILABLE"
    allocation.transport_request.status = "Completed"
    record_transport_event(allocation.transport_request, "Trip completed and vehicle returned", "Completed", f"Distance travelled: {odometer_end - allocation.odometer_start} km.")
    db.session.commit()
    log_audit("Transport", "complete", "transport_allocation", f"Completed allocation {allocation.id}; distance {odometer_end - allocation.odometer_start}.", user=current_user)
    flash("Trip completed and vehicle returned.")
    return redirect(url_for("transport.index"))


@transport_bp.route("/requests/<int:request_id>/documents", methods=["POST"])
@login_required
def upload_transport_document(request_id):
    require_central_transport_role()
    transport_request = TransportPortalRequest.query.get_or_404(request_id)
    upload = request.files.get("document")
    if not upload or not upload.filename:
        flash("Choose a transport case document.")
        return redirect(url_for("transport.transport_case", request_id=request_id)), 400
    original_filename = secure_filename(upload.filename)
    extension = original_filename.rsplit(".", 1)[-1].lower() if "." in original_filename else ""
    if extension not in {"pdf", "doc", "docx", "xls", "xlsx", "csv", "png", "jpg", "jpeg"}:
        flash("Transport documents must be PDF, Word, Excel, CSV, or image files.")
        return redirect(url_for("transport.transport_case", request_id=request_id)), 400
    document_type = request.form.get("document_type", "Supporting document").strip() or "Supporting document"
    prior_versions = TransportDocument.query.filter_by(transport_request_id=request_id, document_type=document_type, original_filename=original_filename).all()
    version = max((document.version for document in prior_versions), default=0) + 1
    folder = os.path.join(current_app.instance_path, "transport_documents", str(request_id))
    os.makedirs(folder, exist_ok=True)
    stored_name = f"{version}_{uuid.uuid4().hex}_{original_filename}"
    upload.save(os.path.join(folder, stored_name))
    db.session.add(TransportDocument(transport_request=transport_request, document_type=document_type, original_filename=original_filename, version=version, storage_path=os.path.join(folder, stored_name), uploaded_by=current_user))
    record_transport_event(transport_request, "Document uploaded", transport_request.status, f"{document_type}: {original_filename} version {version}.")
    db.session.commit()
    log_audit("SmartFleet", "document_upload", "transport_document", f"Uploaded {original_filename} to case {request_id}.", user=current_user)
    flash("Transport document uploaded as a new version.")
    return redirect(url_for("transport.transport_case", request_id=request_id))


@transport_bp.route("/vehicles/<int:vehicle_id>/maintenance", methods=["POST"])
@login_required
def create_transport_maintenance(vehicle_id):
    require_central_transport_role()
    vehicle = Vehicle.query.get_or_404(vehicle_id)
    record = TransportMaintenance(vehicle=vehicle, maintenance_type=request.form.get("maintenance_type", "Service"), description=request.form.get("description"), vendor=request.form.get("vendor"), cost=float(request.form.get("cost", 0) or 0))
    vehicle.availability_status = "MAINTENANCE"
    db.session.add(record)
    db.session.commit()
    log_audit("SmartFleet", "maintenance_opened", "vehicle", f"Vehicle {vehicle.fleet_number} placed into maintenance.", user=current_user)
    flash(f"Maintenance opened for {vehicle.fleet_number}; it is unavailable for allocation.")
    return redirect(url_for("transport.vehicles"))


@transport_bp.route("/vehicles/<int:vehicle_id>/maintenance/<int:maintenance_id>/close", methods=["POST"])
@login_required
def close_transport_maintenance(vehicle_id, maintenance_id):
    require_central_transport_role()
    record = TransportMaintenance.query.filter_by(id=maintenance_id, vehicle_id=vehicle_id).first_or_404()
    record.status = "CLOSED"
    record.completed_at = datetime.utcnow()
    record.vehicle.availability_status = "AVAILABLE"
    db.session.commit()
    log_audit("SmartFleet", "maintenance_closed", "vehicle", f"Maintenance closed for {record.vehicle.fleet_number}.", user=current_user)
    flash(f"Maintenance closed; {record.vehicle.fleet_number} is available.")
    return redirect(url_for("transport.vehicles"))


@transport_bp.route("/requests/<int:request_id>/incidents", methods=["POST"])
@login_required
def create_transport_incident(request_id):
    require_central_transport_role()
    transport_request = TransportPortalRequest.query.get_or_404(request_id)
    allocation = transport_request.allocation
    incident = TransportIncident(transport_request=transport_request, vehicle=allocation.vehicle if allocation else None, driver=allocation.driver if allocation else None, incident_type=request.form.get("incident_type", "Other"), description=request.form.get("description", "").strip(), location=request.form.get("location"), damage=request.form.get("damage"))
    if not incident.description:
        flash("Incident description is required.")
        return redirect(url_for("transport.transport_case", request_id=request_id)), 400
    db.session.add(incident)
    record_transport_event(transport_request, "Incident recorded", transport_request.status, f"{incident.incident_type}: {incident.description}")
    db.session.commit()
    log_audit("SmartFleet", "incident", "transport_incident", f"Incident recorded for case {request_id}.", user=current_user)
    flash("Transport incident recorded for review.")
    return redirect(url_for("transport.transport_case", request_id=request_id))


transport_portal_bp = Blueprint("transport_portal", __name__, url_prefix="/transport-request")


@transport_portal_bp.route("/", methods=["GET", "POST"])
@login_required
def request_transport():
    require_directorate_transport_role()
    employee = current_user.employee
    if request.method == "POST":
        portal_request = TransportPortalRequest(
            requester=current_user,
            employee_name=request.form.get("employee_name") or current_user.full_name,
            employee_number=request.form.get("employee_number") or (employee.employee_number if employee else current_user.username),
            directorate=request.form.get("directorate") or (employee.directorate.name if employee and employee.directorate else None),
            unit=request.form.get("unit") or (employee.unit.name if employee and employee.unit else None),
            contact_details=request.form.get("contact_details") or (employee.phone if employee else current_user.email),
            purpose=request.form.get("purpose"),
            travel_date=request.form.get("travel_date"),
            departure_time=request.form.get("departure_time"),
            return_time=request.form.get("return_time"),
            pickup_location=request.form.get("pickup_location"),
            destination=request.form.get("destination"),
            passenger_count=safe_int(request.form.get("passenger_count"), 1),
            additional_stops=request.form.get("additional_stops"),
            special_requirements=request.form.get("special_requirements"),
            status="Submitted",
        )
        db.session.add(portal_request)
        db.session.flush()
        record_transport_event(portal_request, "Directorate request created", "Draft", "Created by an authorised directorate official.")
        portal_request.status = "Submitted"
        record_transport_event(portal_request, "Submitted to Central Transport SCM", "Submitted", "Directorate submission pack entered the Central SCM queue.")
        db.session.commit()
        log_audit("Transport Portal", "create", "transport_portal_request", f"Submitted transport request {portal_request.id}.", user=current_user)
        flash("Transport request submitted successfully.")
        return redirect(url_for("transport_portal.my_transport_requests"))
    return render_template("transport_portal_form.html", employee=employee)


@transport_portal_bp.route("/mine")
@login_required
def my_transport_requests():
    require_directorate_transport_role()
    if current_user.role_obj.name == "System Administrator":
        items = TransportPortalRequest.query.order_by(TransportPortalRequest.created_at.desc()).all()
    else:
        directorate_name = current_user.employee.directorate.name if current_user.employee and current_user.employee.directorate else None
        items = TransportPortalRequest.query.filter_by(directorate=directorate_name).order_by(TransportPortalRequest.created_at.desc()).all()
    return render_template("transport_portal_list.html", items=items)


travel_bp = Blueprint("travel", __name__, url_prefix="/travel")


@travel_bp.route("/")
@login_required
def index():
    items = TravelRequest.query.order_by(TravelRequest.created_at.desc()).all()
    return render_template("module_list.html", title="Travel / VA26", items=items, columns=["Title", "Destination", "Status"])


@travel_bp.route("/new", methods=["GET", "POST"])
@login_required
def create_travel_request():
    if request.method == "POST":
        travel = TravelRequest(
            title=request.form.get("title"),
            purpose=request.form.get("purpose"),
            destination=request.form.get("destination"),
            status=request.form.get("status", "Draft"),
        )
        db.session.add(travel)
        db.session.commit()
        log_audit("Travel", "create", "travel_request", f"Created travel request {travel.title}.", user=current_user)
        flash("Travel request created successfully.")
        return redirect(url_for("travel.index"))
    return render_template("travel_form.html")


# Assets and controls
asset_bp = Blueprint("asset", __name__, url_prefix="/asset")


@asset_bp.route("/")
@login_required
def index():
    items = Asset.query.order_by(Asset.created_at.desc()).all()
    return render_template("module_list.html", title="Assets", items=items, columns=["Asset Number", "Description", "Condition", "Status"])


@asset_bp.route("/new", methods=["GET", "POST"])
@login_required
def create_asset():
    if request.method == "POST":
        asset = Asset(
            asset_number=request.form.get("asset_number"),
            description=request.form.get("description"),
            category=request.form.get("category"),
            location=request.form.get("location"),
            value=float(request.form.get("value", 0) or 0),
            condition=request.form.get("condition", "Good"),
            status=request.form.get("status", "Active"),
        )
        db.session.add(asset)
        db.session.commit()
        log_audit("Assets", "create", "asset", f"Created asset {asset.asset_number}.", user=current_user)
        flash("Asset created successfully.")
        return redirect(url_for("asset.index"))
    return render_template("asset_form.html")


control_bp = Blueprint("control", __name__, url_prefix="/controls")


@control_bp.route("/")
@login_required
def index():
    items = InternalControl.query.order_by(InternalControl.created_at.desc()).all()
    return render_template("module_list.html", title="Internal Controls", items=items, columns=["Control Objective", "Process", "Risk Rating", "Status"])


@control_bp.route("/new", methods=["GET", "POST"])
@login_required
def create_internal_control():
    if request.method == "POST":
        control = InternalControl(
            control_objective=request.form.get("control_objective"),
            process_name=request.form.get("process_name"),
            risk_rating=request.form.get("risk_rating", "Medium"),
            status=request.form.get("status", "Active"),
        )
        db.session.add(control)
        db.session.commit()
        log_audit("Internal Controls", "create", "internal_control", f"Created control {control.control_objective}.", user=current_user)
        flash("Internal control created successfully.")
        return redirect(url_for("control.index"))
    return render_template("control_form.html")


risk_bp = Blueprint("risk", __name__, url_prefix="/risk")


@risk_bp.route("/")
@login_required
def index():
    items = RiskRegister.query.order_by(RiskRegister.created_at.desc()).all()
    return render_template("module_list.html", title="Risk Management", items=items, columns=["Title", "Category", "Risk Rating", "Status"])


@risk_bp.route("/new", methods=["GET", "POST"])
@login_required
def create_risk_item():
    if request.method == "POST":
        risk = RiskRegister(
            title=request.form.get("title"),
            category=request.form.get("category"),
            likelihood=request.form.get("likelihood", "Medium"),
            impact=request.form.get("impact", "Medium"),
            risk_rating=request.form.get("risk_rating", "Moderate"),
            status=request.form.get("status", "Open"),
        )
        db.session.add(risk)
        db.session.commit()
        log_audit("Risk Management", "create", "risk_register", f"Created risk {risk.title}.", user=current_user)
        flash("Risk item created successfully.")
        return redirect(url_for("risk.index"))
    return render_template("risk_form.html")


# Audit, reports, security

audit_bp = Blueprint("audit", __name__, url_prefix="/audit")


@audit_bp.route("/")
@login_required
def index():
    items = AuditLog.query.order_by(AuditLog.created_at.desc()).limit(20).all()
    return render_template("module_list.html", title="Audit Log", items=items, columns=["Module", "Action", "Entity", "Created"])


report_bp = Blueprint("reports", __name__, url_prefix="/reports")


@report_bp.route("/")
@login_required
def index():
    return render_template("report_index.html")


security_bp = Blueprint("security", __name__, url_prefix="/security")


@security_bp.route("/")
@login_required
def index():
    items = SecurityEvent.query.order_by(SecurityEvent.created_at.desc()).all()
    return render_template("module_list.html", title="SmartShield", items=items, columns=["Severity", "Module", "Status", "Created"])


@security_bp.route("/new", methods=["GET", "POST"])
@login_required
def create_security_event():
    if request.method == "POST":
        event = SecurityEvent(
            severity=request.form.get("severity", "Medium"),
            module=request.form.get("module", "General"),
            description=request.form.get("description"),
            status=request.form.get("status", "Open"),
        )
        db.session.add(event)
        db.session.commit()
        log_audit("SmartShield", "create", "security_event", f"Created security event {event.module}.", user=current_user)
        flash("Security event created successfully.")
        return redirect(url_for("security.index"))
    return render_template("security_form.html")


# Documents, notifications, search, admin

document_bp = Blueprint("document", __name__, url_prefix="/documents")


@document_bp.route("/")
@login_required
def index():
    items = Document.query.order_by(Document.created_at.desc()).all()
    return render_template("module_list.html", title="Document Management", items=items, columns=["Title", "Type", "Uploaded By", "Date"])


@document_bp.route("/new", methods=["GET", "POST"])
@login_required
def create_document():
    if request.method == "POST":
        document = Document(
            title=request.form.get("title"),
            document_type=request.form.get("document_type", "General"),
            link_type=request.form.get("link_type"),
            link_id=safe_int(request.form.get("link_id")),
            uploaded_by=request.form.get("uploaded_by") or current_user.full_name,
        )
        db.session.add(document)
        db.session.commit()
        log_audit("Documents", "create", "document", f"Created document {document.title}.", user=current_user)
        flash("Document created successfully.")
        return redirect(url_for("document.index"))
    return render_template("document_form.html")


notification_bp = Blueprint("notifications", __name__, url_prefix="/notifications")


@notification_bp.route("/")
@login_required
def index():
    items = Notification.query.filter_by(user_id=current_user.id).order_by(Notification.created_at.desc()).all()
    return render_template("module_list.html", title="Notifications", items=items, columns=["Title", "Message", "Read"])


@notification_bp.route("/new", methods=["GET", "POST"])
@login_required
def create_notification():
    if request.method == "POST":
        notification = Notification(
            user_id=current_user.id,
            title=request.form.get("title"),
            message=request.form.get("message"),
            is_read=request.form.get("is_read") == "true",
        )
        db.session.add(notification)
        db.session.commit()
        log_audit("Notifications", "create", "notification", f"Created notification {notification.title}.", user=current_user)
        flash("Notification created successfully.")
        return redirect(url_for("notifications.index"))
    return render_template("notification_form.html")


search_bp = Blueprint("search", __name__, url_prefix="/search")


@search_bp.route("/")
@login_required
def index():
    query = request.args.get("q", "")
    results = []
    if query:
        results = Employee.query.filter(Employee.first_name.ilike(f"%{query}%") | Employee.last_name.ilike(f"%{query}%")).all()
    return render_template("search.html", query=query, results=results)


admin_bp = Blueprint("admin", __name__, url_prefix="/admin")


@admin_bp.route("/")
@login_required
def index():
    roles = Role.query.order_by(Role.name).all()
    return render_template("admin.html", roles=roles)


@admin_bp.route("/backups")
@login_required
def backups():
    items = BackupRecord.query.order_by(BackupRecord.created_at.desc()).all()
    return render_template("backup_records.html", items=items)


approval_bp = Blueprint("approvals", __name__, url_prefix="/approvals")


APPROVAL_ROLES = {"System Administrator", "SCM Manager", "CFO / Finance Authority", "DSD Executive / Executive Management", "Internal Control Officer"}


@approval_bp.route("/")
@login_required
def index():
    status = request.args.get("status")
    query = ApprovalRecord.query
    if status:
        query = query.filter_by(status=status)
    if status and status.lower() == "pending" and current_user.role_obj:
        permission_names = {permission.name for permission in current_user.role_obj.permissions}
        can_review_all = current_user.role_obj.name == "System Administrator" or bool(permission_names.intersection({"approve_requisition", "approve_procurement", "approve_budget"}))
        if not can_review_all:
            query = query.filter(ApprovalRecord.approver.in_([current_user.username, current_user.full_name]))
    items = query.order_by(ApprovalRecord.created_at.desc()).all()
    empty_message = "You have no pending approvals." if status and status.lower() == "pending" else "No approval records here yet"
    return render_template("module_list.html", title="Approvals", items=items, columns=["Title", "Entity Type", "Approver", "Status"], active_filter=status, empty_message=empty_message)


@approval_bp.route("/new", methods=["GET", "POST"])
@login_required
def create_approval():
    if request.method == "POST":
        approval = ApprovalRecord(
            title=request.form.get("title"),
            entity_type=request.form.get("entity_type", "General"),
            entity_id=safe_int(request.form.get("entity_id"), 1),
            approver=request.form.get("approver"),
            status=request.form.get("status", "Pending"),
            comments=request.form.get("comments"),
        )
        db.session.add(approval)
        db.session.commit()
        log_audit("Approvals", "create", "approval_record", f"Created approval {approval.title}.", user=current_user)
        flash("Approval record created successfully.")
        return redirect(url_for("approvals.index"))
    return render_template("approval_form.html")


@approval_bp.route("/<int:approval_id>/decision", methods=["POST"])
@login_required
def decide_approval(approval_id):
    if not current_user.role_obj or current_user.role_obj.name not in APPROVAL_ROLES:
        abort(403)
    approval = ApprovalRecord.query.get_or_404(approval_id)
    if approval.status != "Pending":
        flash("Only pending approvals can receive a decision.")
        return redirect(url_for("approvals.index")), 409
    decision = request.form.get("decision")
    if decision not in {"Approved", "Rejected", "Returned", "Correction Required"}:
        flash("Select a valid approval decision.")
        return redirect(url_for("approvals.index")), 400
    approval.status = decision
    approval.comments = request.form.get("comments")
    db.session.commit()
    log_audit("Approvals", decision.lower().replace(" ", "_"), "approval_record", f"Approval {approval.id} decided as {decision}.", user=current_user)
    flash(f"Approval {decision.lower()}.")
    return redirect(url_for("approvals.index"))


policy_bp = Blueprint("policies", __name__, url_prefix="/policies")


@policy_bp.route("/")
@login_required
def index():
    items = Policy.query.order_by(Policy.created_at.desc()).all()
    return render_template("module_list.html", title="Policies", items=items, columns=["Title", "Category", "Version", "Status"])


@policy_bp.route("/new", methods=["GET", "POST"])
@login_required
def create_policy():
    if request.method == "POST":
        policy = Policy(
            title=request.form.get("title"),
            category=request.form.get("category", "General"),
            version=request.form.get("version"),
            owner=request.form.get("owner"),
            status=request.form.get("status", "Draft"),
        )
        db.session.add(policy)
        db.session.commit()
        log_audit("Policies", "create", "policy", f"Created policy {policy.title}.", user=current_user)
        flash("Policy created successfully.")
        return redirect(url_for("policies.index"))
    return render_template("policy_form.html")


archive_bp = Blueprint("archive", __name__, url_prefix="/archive")


@archive_bp.route("/")
@login_required
def index():
    items = ArchiveRecord.query.order_by(ArchiveRecord.created_at.desc()).all()
    return render_template("module_list.html", title="Archive", items=items, columns=["Title", "Type", "Status"])


@archive_bp.route("/new", methods=["GET", "POST"])
@login_required
def create_archive_record():
    if request.method == "POST":
        archive = ArchiveRecord(
            title=request.form.get("title"),
            archive_type=request.form.get("archive_type", "General"),
            storage_path=request.form.get("storage_path"),
            status=request.form.get("status", "Archived"),
        )
        db.session.add(archive)
        db.session.commit()
        log_audit("Archive", "create", "archive_record", f"Created archive {archive.title}.", user=current_user)
        flash("Archive record created successfully.")
        return redirect(url_for("archive.index"))
    return render_template("archive_form.html")


performance_bp = Blueprint("performance", __name__, url_prefix="/performance")


@performance_bp.route("/")
@login_required
def index():
    items = PerformanceIndicator.query.order_by(PerformanceIndicator.created_at.desc()).all()
    return render_template("module_list.html", title="Performance Indicators", items=items, columns=["Title", "Category", "Target", "Actual", "Status"])


@performance_bp.route("/new", methods=["GET", "POST"])
@login_required
def create_performance_indicator():
    if request.method == "POST":
        indicator = PerformanceIndicator(
            title=request.form.get("title"),
            category=request.form.get("category", "General"),
            target=request.form.get("target"),
            actual=request.form.get("actual"),
            status=request.form.get("status", "On Track"),
        )
        db.session.add(indicator)
        db.session.commit()
        log_audit("Performance", "create", "performance_indicator", f"Created performance indicator {indicator.title}.", user=current_user)
        flash("Performance indicator created successfully.")
        return redirect(url_for("performance.index"))
    return render_template("performance_form.html")


compliance_bp = Blueprint("compliance", __name__, url_prefix="/compliance")


@compliance_bp.route("/")
@login_required
def index():
    items = ComplianceCheck.query.order_by(ComplianceCheck.created_at.desc()).all()
    return render_template("module_list.html", title="Compliance Checks", items=items, columns=["Title", "Category", "Owner", "Status"])


@compliance_bp.route("/new", methods=["GET", "POST"])
@login_required
def create_compliance_check():
    if request.method == "POST":
        check = ComplianceCheck(
            title=request.form.get("title"),
            category=request.form.get("category", "General"),
            owner=request.form.get("owner"),
            status=request.form.get("status", "Monitoring"),
            notes=request.form.get("notes"),
        )
        db.session.add(check)
        db.session.commit()
        log_audit("Compliance", "create", "compliance_check", f"Created compliance check {check.title}.", user=current_user)
        flash("Compliance check created successfully.")
        return redirect(url_for("compliance.index"))
    return render_template("compliance_form.html")


budget_bp = Blueprint("budget", __name__, url_prefix="/budget")


@budget_bp.route("/")
@login_required
def index():
    items = BudgetForecast.query.order_by(BudgetForecast.created_at.desc()).all()
    return render_template("module_list.html", title="Budget Forecasts", items=items, columns=["Year", "Department", "Allocation", "Expenditure", "Variance", "Status"])


@budget_bp.route("/new", methods=["GET", "POST"])
@login_required
def create_budget_forecast():
    if request.method == "POST":
        budget = BudgetForecast(
            fiscal_year=request.form.get("fiscal_year"),
            department=request.form.get("department"),
            allocation=float(request.form.get("allocation", 0) or 0),
            expenditure=float(request.form.get("expenditure", 0) or 0),
            variance=float(request.form.get("variance", 0) or 0),
            status=request.form.get("status", "On Track"),
        )
        db.session.add(budget)
        db.session.commit()
        log_audit("Budget", "create", "budget_forecast", f"Created budget forecast for {budget.department}.", user=current_user)
        flash("Budget forecast created successfully.")
        return redirect(url_for("budget.index"))
    return render_template("budget_form.html")


portfolio_bp = Blueprint("portfolio", __name__, url_prefix="/portfolio")


@portfolio_bp.route("/")
@login_required
def index():
    items = PortfolioProject.query.order_by(PortfolioProject.created_at.desc()).all()
    return render_template("module_list.html", title="Portfolio Projects", items=items, columns=["Project", "Phase", "Owner", "Due Date", "Progress", "Status"])


@portfolio_bp.route("/new", methods=["GET", "POST"])
@login_required
def create_portfolio_project():
    if request.method == "POST":
        project = PortfolioProject(
            project_name=request.form.get("project_name"),
            phase=request.form.get("phase", "Planning"),
            owner=request.form.get("owner"),
            due_date=request.form.get("due_date"),
            completion_percent=safe_int(request.form.get("completion_percent"), 0),
            status=request.form.get("status", "Planned"),
        )
        db.session.add(project)
        db.session.commit()
        log_audit("Portfolio", "create", "portfolio_project", f"Created portfolio project {project.project_name}.", user=current_user)
        flash("Portfolio project created successfully.")
        return redirect(url_for("portfolio.index"))
    return render_template("portfolio_form.html")


strategic_plans_bp = Blueprint("strategic_plans", __name__, url_prefix="/strategic-plans")


@strategic_plans_bp.route("/")
@login_required
def index():
    items = StrategicPlan.query.order_by(StrategicPlan.created_at.desc()).all()
    return render_template("module_list.html", title="Strategic Plans", items=items, columns=["Objective", "Outcome", "Owner", "Target Date", "Status", "Progress"])


@strategic_plans_bp.route("/new", methods=["GET", "POST"])
@login_required
def create_strategic_plan():
    if request.method == "POST":
        plan = StrategicPlan(
            objective=request.form.get("objective"),
            outcome=request.form.get("outcome"),
            owner=request.form.get("owner"),
            target_date=request.form.get("target_date"),
            status=request.form.get("status", "Planned"),
            progress_percent=safe_int(request.form.get("progress_percent"), 0),
            notes=request.form.get("notes"),
        )
        db.session.add(plan)
        db.session.commit()
        log_audit("Strategic Planning", "create", "strategic_plan", f"Created strategic objective {plan.objective}.", user=current_user)
        flash("Strategic plan created successfully.")
        return redirect(url_for("strategic_plans.index"))
    return render_template("strategic_plan_form.html")


stakeholders_bp = Blueprint("stakeholders", __name__, url_prefix="/stakeholders")


@stakeholders_bp.route("/")
@login_required
def index():
    items = Stakeholder.query.order_by(Stakeholder.created_at.desc()).all()
    return render_template("module_list.html", title="Stakeholders", items=items, columns=["Name", "Organisation", "Role", "Influence", "Status"])


@stakeholders_bp.route("/new", methods=["GET", "POST"])
@login_required
def create_stakeholder():
    if request.method == "POST":
        stakeholder = Stakeholder(
            name=request.form.get("name"),
            organisation=request.form.get("organisation"),
            role=request.form.get("role"),
            influence=request.form.get("influence", "Medium"),
            engagement_status=request.form.get("engagement_status", "Active"),
            notes=request.form.get("notes"),
        )
        db.session.add(stakeholder)
        db.session.commit()
        log_audit("Stakeholders", "create", "stakeholder", f"Created stakeholder {stakeholder.name}.", user=current_user)
        flash("Stakeholder created successfully.")
        return redirect(url_for("stakeholders.index"))
    return render_template("stakeholder_form.html")


training_bp = Blueprint("training", __name__, url_prefix="/training")


@training_bp.route("/")
@login_required
def index():
    items = TrainingRecord.query.order_by(TrainingRecord.created_at.desc()).all()
    return render_template("module_list.html", title="Training Records", items=items, columns=["Title", "Category", "Provider", "Target Group", "Status"])


@training_bp.route("/new", methods=["GET", "POST"])
@login_required
def create_training_record():
    if request.method == "POST":
        training = TrainingRecord(
            title=request.form.get("title"),
            category=request.form.get("category", "General"),
            provider=request.form.get("provider"),
            target_group=request.form.get("target_group"),
            completion_date=request.form.get("completion_date"),
            status=request.form.get("status", "Scheduled"),
        )
        db.session.add(training)
        db.session.commit()
        log_audit("Training", "create", "training_record", f"Created training record {training.title}.", user=current_user)
        flash("Training record created successfully.")
        return redirect(url_for("training.index"))
    return render_template("training_form.html")


beneficiaries_bp = Blueprint("beneficiaries", __name__, url_prefix="/beneficiaries")


@beneficiaries_bp.route("/")
@login_required
def index():
    items = Beneficiary.query.order_by(Beneficiary.created_at.desc()).all()
    return render_template("module_list.html", title="Beneficiaries", items=items, columns=["Name", "Program", "Location", "Priority", "Status"])


@beneficiaries_bp.route("/new", methods=["GET", "POST"])
@login_required
def create_beneficiary():
    if request.method == "POST":
        beneficiary = Beneficiary(
            beneficiary_name=request.form.get("beneficiary_name"),
            program=request.form.get("program"),
            location=request.form.get("location"),
            priority=request.form.get("priority", "Medium"),
            status=request.form.get("status", "Active"),
            notes=request.form.get("notes"),
        )
        db.session.add(beneficiary)
        db.session.commit()
        log_audit("Beneficiaries", "create", "beneficiary", f"Created beneficiary {beneficiary.beneficiary_name}.", user=current_user)
        flash("Beneficiary created successfully.")
        return redirect(url_for("beneficiaries.index"))
    return render_template("beneficiary_form.html")


service_feedback_bp = Blueprint("service_feedback", __name__, url_prefix="/service-feedback")


@service_feedback_bp.route("/")
@login_required
def index():
    items = ServiceFeedback.query.order_by(ServiceFeedback.created_at.desc()).all()
    return render_template("module_list.html", title="Service Feedback", items=items, columns=["Service", "Beneficiary", "Rating", "Status"])


@service_feedback_bp.route("/new", methods=["GET", "POST"])
@login_required
def create_service_feedback():
    if request.method == "POST":
        feedback = ServiceFeedback(
            service_name=request.form.get("service_name"),
            beneficiary_name=request.form.get("beneficiary_name"),
            rating=safe_int(request.form.get("rating"), 0),
            comment=request.form.get("comment"),
            status=request.form.get("status", "Open"),
        )
        db.session.add(feedback)
        db.session.commit()
        log_audit("Service Feedback", "create", "service_feedback", f"Created service feedback for {feedback.beneficiary_name}.", user=current_user)
        flash("Service feedback created successfully.")
        return redirect(url_for("service_feedback.index"))
    return render_template("service_feedback_form.html")


@admin_bp.route("/health")
def health():
    return {"status": "ok", "app": current_app.config["APP_NAME"], "version": current_app.config["APP_VERSION"]}
